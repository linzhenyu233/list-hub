# -*- coding: utf-8 -*-
"""批量商品中台首期服务：Excel 导入、校验、批次查询与模板下载。

本服务刻意独立于微信/小红书服务。当前环境没有 MySQL 驱动和 multipart，
因此上传接口采用 JSON + Base64，存储使用 SQLite；后续可替换存储层而不改前端接口。
"""
import base64
import csv
import io
import hashlib
import json
import os
import re
import shutil
import sqlite3
import urllib.parse
import urllib.request
import urllib.error
import uuid
import zipfile
import time
import sys
from datetime import datetime, timezone

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from runtime_config import load_project_env
import shop_registry

load_project_env()


# ---- 店铺上下文（请求级） ------------------------------------------------
# 路由层通过 _current_shop(request) 拿到当前店铺；未传 X-Shop-Id 则回退默认店。
# 操作人通过 X-Operator 头获取，空则取默认。

SHOP_HEADER = "x-shop-id"
OPERATOR_HEADER = "x-operator"


def _current_shop_id(request: Request) -> str:
    """从请求头取 shop_id，空则回退默认店。未知/停用店直接抛 400。"""
    sid = request.headers.get(SHOP_HEADER) or ""
    sid = sid.strip()
    if not sid:
        return shop_registry.default_shop()["shop_id"]
    shop = _verify_shop(sid)
    return shop["shop_id"]


def _current_operator(request: Request) -> str:
    """从请求头取操作人，空则取默认。

    前端会做 encodeURIComponent：HTTP 头只允许 ASCII，中文姓名直接放会报
    "Cannot convert argument to a ByteString"。这里 decode 回来；
    若调用方（脚本/curl）没编码，unquote 对纯 ASCII 是空操作，向后兼容。
    """
    op = request.headers.get(OPERATOR_HEADER) or ""
    return urllib.parse.unquote(op.strip()) or shop_registry.default_operator()


def _verify_shop(shop_id: str) -> dict:
    """校验店铺存在且启用，返回店铺配置。不存在/停用抛 HTTPException(400)。"""
    try:
        return shop_registry.resolve(shop_id)
    except ValueError as e:
        raise HTTPException(400, str(e))

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.environ.get("BULK_DB_FILE", os.path.join(BASE_DIR, "bulk_catalog.sqlite3"))
IMAGE_DIR = os.environ.get("BULK_IMAGE_DIR", os.path.join(BASE_DIR, "../images"))
BULK_API_HOST = os.environ.get("BULK_API_HOST", "0.0.0.0")
BULK_API_PORT = int(os.environ.get("BULK_API_PORT", "8020"))
os.makedirs(IMAGE_DIR, exist_ok=True)

# 图片根目录默认值(网络共享, 服务端可直读):图片直读扫描的默认根目录
DEFAULT_IMAGE_ROOT = r"\\192.168.10.250\电子商务部\网销部共享\SHINING HOUSE培育钻"
# 允许按引用读取图片的根目录白名单:
# 防止有人构造 disk:/local: 引用(或直接传 root 参数)去读服务器上的任意文件
ALLOWED_IMAGE_ROOTS = [DEFAULT_IMAGE_ROOT, IMAGE_DIR]


def _within_allowed_roots(path, roots=None):
    """判断绝对路径是否落在白名单根目录内。默认用全局 ALLOWED_IMAGE_ROOTS。

    用 commonpath 而不是 startswith:后者会让 '...\\images_evil\\x.jpg' 命中
    白名单根 '...\\images' 的前缀,导致越权读取。
    """
    if roots is None:
        roots = ALLOWED_IMAGE_ROOTS
    target = os.path.abspath(str(path or ""))
    for root in roots:
        root_abs = os.path.abspath(root)
        try:
            if os.path.commonpath((root_abs, target)) == root_abs:
                return True
        except ValueError:
            # 不同盘符 / 本地盘与 UNC 混用时 commonpath 会报错,视为不在白名单内
            continue
    return False


def _allowed_roots_for(shop_id=None):
    """当前店铺允许读图的根目录白名单：店铺 image_root + 全局兜底。

    统一入口 —— 原先 4 个校验点各写一遍，其中发布读图那处漏了店铺 root，
    导致「扫描时认可、发布时拒绝」：服务器 image_root 是 COS 挂载点 /mnt/cos/...，
    而 DEFAULT_IMAGE_ROOT 是为 Windows 共享盘写的 UNC 路径，两者永不相等。
    """
    shop_root = ""
    if shop_id:
        try:
            shop_root = shop_registry.image_root_for(shop_id)
        except ValueError:
            shop_root = ""  # 店铺已停用/不存在：退回全局兜底，由调用方给出明确报错
    return list(dict.fromkeys(([shop_root] if shop_root else []) + list(ALLOWED_IMAGE_ROOTS)))


app = FastAPI(title="商品批量发布中台", version="0.1.0")
_WECHAT_CATEGORY_CHAIN_CACHE = {}
_XHS_VAR_CANDIDATES_CACHE = {}

# 图片上传缓存:内存 LRU 层(SQLite 之上),命中直接返回,省一次磁盘查询。
# 上限 2000 条(按单商品约 10 张图算,覆盖 200 件商品的并发上传窗口)。
_IMAGE_CACHE_MEM = {}
_IMAGE_CACHE_ORDER = []  # 插入顺序,用于 LRU 淘汰
_IMAGE_CACHE_MAX = 2000
_image_cache_hit_mem = 0  # 内存命中计数
_image_cache_hit_db = 0   # DB 命中计数
_image_cache_miss = 0     # 未命中计数

# 小红书 SKU 规格值别名:货盘英文 SKU 编码 → 平台中文 valueName(发品时映射用)
# 小红书「颜色分类」等规格的 valueName 是中文,货盘常用英文编码,直接发英文会显示英文
XHS_SPEC_VALUE_ALIASES = {
    '粉红色': ['SHINING PINK', '粉色', '粉红', 'PINK', 'Rose'],
    '浅蓝色': ['ICE BLUE', '浅蓝', 'LIGHT BLUE'],
    '香槟金色': ['CHAMPAGNE GOLD', '香槟金', 'CHAMPAGNE'],
    '玫瑰金色': ['ROSE GOLD', '玫瑰金'],
    '黄金色': ['YELLOW GOLD', '黄金', 'GOLD'],
    '18K金色': ['18K GOLD', '18K金'],
    '白色': ['WHITE'],
    '黑色': ['BLACK'],
    '无色': ['COLORLESS', '透明'],
    '玫红色': ['ROSE RED', '玫红'],
    '红色': ['RED'],
    '绿色': ['GREEN'],
    '黄色': ['YELLOW'],
    '蓝色': ['BLUE'],
    '紫色': ['PURPLE', 'VIOLET'],
    '灰色': ['GRAY', 'GREY'],
    '银色': ['SILVER'],
    # 其它颜色/规格值按需补充
}


def db():
    conn = sqlite3.connect(DB_FILE, timeout=30)
    conn.row_factory = sqlite3.Row
    # WAL 模式:读写互不阻塞,大幅降低 database is locked
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=30000")
    conn.execute("PRAGMA synchronous=NORMAL")

    # 顺序很重要：先建表（新库直接是新结构）→ 补列 → 重建带单店约束的老表
    # → 迁移历史数据归属 → 最后建索引。
    # 如果索引先建，老库还没有 shop_id 列，CREATE INDEX 会直接报 no such column。
    _create_tables(conn)
    _migrate_add_shop_id(conn)
    _migrate_add_job_params(conn)
    _rebuild_multishop_tables(conn)
    _reassign_legacy_shop(conn)
    _create_indexes(conn)

    conn.commit()
    return conn


def _create_tables(conn):
    """建表（仅 CREATE TABLE IF NOT EXISTS，不含索引）。"""
    conn.execute("CREATE TABLE IF NOT EXISTS _meta (key TEXT PRIMARY KEY, value TEXT)")
    conn.execute("""CREATE TABLE IF NOT EXISTS batches (
        id TEXT PRIMARY KEY, shop_id TEXT NOT NULL, filename TEXT, status TEXT,
        total INTEGER DEFAULT 0, valid INTEGER DEFAULT 0, errors INTEGER DEFAULT 0,
        created_at TEXT, operator TEXT,
        -- 每次 rows_json / mappings_json 被改写就 +1: 供跨进程批次缓存判新旧
        revision INTEGER NOT NULL DEFAULT 0,
        rows_json TEXT NOT NULL, mappings_json TEXT NOT NULL DEFAULT '{}'
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS category_aliases (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        shop_id TEXT NOT NULL,
        internal_category TEXT NOT NULL,
        wechat_json TEXT, xhs_json TEXT, updated_at TEXT,
        UNIQUE(shop_id, internal_category)
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS publish_jobs (
        id TEXT PRIMARY KEY, shop_id TEXT NOT NULL, batch_id TEXT NOT NULL,
        platforms_json TEXT NOT NULL, status TEXT NOT NULL, total INTEGER DEFAULT 0,
        processed INTEGER DEFAULT 0, success INTEGER DEFAULT 0, failed INTEGER DEFAULT 0,
        operator TEXT, created_at TEXT, updated_at TEXT,
        -- 跨店发布：目标店铺的私有参数快照(运费模板/物流方案/品牌)。
        -- shop_id 即目标店；本字段让 worker 执行时覆盖批次里的来源店参数。空对象=同店发布。
        params_json TEXT NOT NULL DEFAULT '{}'
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS publish_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT NOT NULL, batch_id TEXT NOT NULL,
        shop_id TEXT NOT NULL, product_code TEXT NOT NULL, platform TEXT NOT NULL,
        status TEXT NOT NULL, platform_product_id TEXT, platform_sku_ids TEXT, error TEXT,
        started_at TEXT, finished_at TEXT, priority INTEGER NOT NULL DEFAULT 0,
        UNIQUE(job_id, product_code, platform)
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS image_upload_cache (
        shop_id TEXT NOT NULL, platform TEXT NOT NULL, source_hash TEXT NOT NULL,
        source_url TEXT NOT NULL, platform_url TEXT NOT NULL, updated_at TEXT,
        PRIMARY KEY(shop_id, platform, source_hash)
    )""")


def _create_indexes(conn):
    """建索引。必须在补列 + 重建表之后调用（老库此时才有 shop_id 列、表约束才是新的）。"""
    conn.execute("CREATE INDEX IF NOT EXISTS idx_batches_shop ON batches(shop_id, created_at DESC)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_category_aliases_shop ON category_aliases(shop_id, updated_at DESC)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_publish_jobs_shop ON publish_jobs(shop_id, created_at DESC)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_publish_items_batch_product_platform ON publish_items(batch_id, product_code, platform)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_publish_items_pick ON publish_items(shop_id, platform, status, priority DESC, id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_publish_items_product ON publish_items(shop_id, product_code, platform, status)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_image_cache_shop ON image_upload_cache(shop_id, platform, updated_at DESC)")
    # 唯一索引（老库靠这两条兜底保证含 shop_id 的唯一性；新库表定义里已内置同名约束）
    conn.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_category_aliases_shop_internal "
        "ON category_aliases(shop_id, internal_category)"
    )
    conn.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_image_cache_shop_platform_hash "
        "ON image_upload_cache(shop_id, platform, source_hash)"
    )


def _migrate_add_shop_id(conn):
    """把老库升级到多店铺版本（幂等）。

    本函数只负责「补列 + 回填」；表级唯一约束的重建在 _rebuild_multishop_tables。

    迁移内容：
      1. priority 列（更早版本遗留）—— 必须在建索引前补齐。
      2. batches.revision 列 —— 供跨进程批次缓存判新旧（见 get_batch_cached）。
      3. 5 张表的 shop_id / operator 列，历史数据回填默认 shop_id。

    默认 shop_id 从 shop_registry.default_shop() 取：
      - 已有 shops.json → 取其默认店铺的 shop_id（推荐：先建好 shops.json 再首次启动）
      - 无 shops.json    → 回退 "default"
    """
    # --- 1. 更早版本遗留的 priority 列（独立于 shop_id，先补） ---
    if "priority" not in {row["name"] for row in conn.execute("PRAGMA table_info(publish_items)")}:
        conn.execute("ALTER TABLE publish_items ADD COLUMN priority INTEGER NOT NULL DEFAULT 0")

    # --- 2. batches.revision：跨进程批次缓存判新旧（见 get_batch_cached）。
    #     必须在下面「已是新版就 return」之前判断，否则已有库永远补不上这一列。 ---
    if "revision" not in {row["name"] for row in conn.execute("PRAGMA table_info(batches)")}:
        conn.execute("ALTER TABLE batches ADD COLUMN revision INTEGER NOT NULL DEFAULT 0")

    # --- 3. shop_id 相关列 ---
    cols = {row["name"] for row in conn.execute("PRAGMA table_info(batches)")}
    if "shop_id" in cols:
        return  # 已是新版

    try:
        default_sid = shop_registry.default_shop()["shop_id"]
    except ValueError:
        default_sid = "default"

    print(f"[db-migrate] 检测到老库，历史数据回填 shop_id = {default_sid}")

    _safe_add_column(conn, "batches", "shop_id TEXT NOT NULL DEFAULT '" + default_sid + "'")
    _safe_add_column(conn, "batches", "operator TEXT")
    _safe_add_column(conn, "category_aliases", "shop_id TEXT NOT NULL DEFAULT '" + default_sid + "'")
    _safe_add_column(conn, "publish_jobs", "shop_id TEXT NOT NULL DEFAULT '" + default_sid + "'")
    _safe_add_column(conn, "publish_jobs", "operator TEXT")
    _safe_add_column(conn, "publish_items", "shop_id TEXT NOT NULL DEFAULT '" + default_sid + "'")
    _safe_add_column(conn, "image_upload_cache", "shop_id TEXT NOT NULL DEFAULT '" + default_sid + "'")

    # 记住这次回填用了哪个 shop_id，供 _reassign_legacy_shop 在配置补齐后自动改名
    _meta_set(conn, "legacy_shop_id", default_sid)

    print("[db-migrate] 补列完成（唯一约束由 _rebuild_multishop_tables 处理）")


def _table_sql(conn, table):
    """取建表 DDL，表不存在返回 None。"""
    try:
        row = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name=?", (table,)
        ).fetchone()
    except sqlite3.OperationalError:
        return None
    return row["sql"] if row and row["sql"] else None


def _rebuild_multishop_tables(conn):
    """重建仍带「单店时代唯一约束」的表。

    单店版本的这两张表有全局唯一约束，ALTER TABLE 加 shop_id 列改不掉它们，
    会直接破坏多店铺语义，必须整表重建（SQLite 不支持 DROP CONSTRAINT）：

      - category_aliases: 旧 UNIQUE(internal_category)
          → 不重建：B 店保存与 A 店同名的内部类目会报
            "UNIQUE constraint failed: category_aliases.internal_category"
            （因为 ON CONFLICT(shop_id, internal_category) 匹配不上旧约束）
          → 新：UNIQUE(shop_id, internal_category)

      - image_upload_cache: 旧 PRIMARY KEY(platform, source_hash)
          → 不重建：B 店上传同一张图会 REPLACE 掉 A 店的缓存行，
            之后 B 店会拿到 A 店素材空间里的 URL（跨店串素材，发错店铺）
          → 新：PRIMARY KEY(shop_id, platform, source_hash)

    幂等：按 DDL 文本判断，已是新结构则跳过；新库由 _create_tables 直接建成新结构。
    """
    # --- category_aliases ---
    sql = _table_sql(conn, "category_aliases")
    if sql:
        normalized = "".join(sql.split()).lower()
        if "unique(shop_id,internal_category)" not in normalized:
            count = conn.execute("SELECT COUNT(*) FROM category_aliases").fetchone()[0]
            print(f"[db-migrate] 重建 category_aliases（{count} 行）以移除单店唯一约束")
            conn.execute("ALTER TABLE category_aliases RENAME TO category_aliases_legacy")
            conn.execute("""CREATE TABLE category_aliases (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                shop_id TEXT NOT NULL,
                internal_category TEXT NOT NULL,
                wechat_json TEXT, xhs_json TEXT, updated_at TEXT,
                UNIQUE(shop_id, internal_category)
            )""")
            conn.execute("""INSERT INTO category_aliases(id,shop_id,internal_category,wechat_json,xhs_json,updated_at)
                SELECT id,shop_id,internal_category,wechat_json,xhs_json,updated_at
                FROM category_aliases_legacy""")
            conn.execute("DROP TABLE category_aliases_legacy")

    # --- image_upload_cache ---
    sql = _table_sql(conn, "image_upload_cache")
    if sql:
        normalized = "".join(sql.split()).lower()
        if "primarykey(shop_id,platform,source_hash)" not in normalized:
            count = conn.execute("SELECT COUNT(*) FROM image_upload_cache").fetchone()[0]
            print(f"[db-migrate] 重建 image_upload_cache（{count} 行）以移除跨店主键冲突")
            conn.execute("ALTER TABLE image_upload_cache RENAME TO image_upload_cache_legacy")
            conn.execute("""CREATE TABLE image_upload_cache (
                shop_id TEXT NOT NULL, platform TEXT NOT NULL, source_hash TEXT NOT NULL,
                source_url TEXT NOT NULL, platform_url TEXT NOT NULL, updated_at TEXT,
                PRIMARY KEY(shop_id, platform, source_hash)
            )""")
            conn.execute("""INSERT OR REPLACE INTO image_upload_cache(shop_id,platform,source_hash,source_url,platform_url,updated_at)
                SELECT shop_id,platform,source_hash,source_url,platform_url,updated_at
                FROM image_upload_cache_legacy""")
            conn.execute("DROP TABLE image_upload_cache_legacy")


def _safe_add_column(conn, table, col_def):
    """执行 ALTER TABLE ADD COLUMN，忽略已存在的列错误。"""
    try:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {col_def}")
    except sqlite3.OperationalError as e:
        if "duplicate column name" in str(e).lower():
            return
        raise


def _migrate_add_job_params(conn):
    """给 publish_jobs 补 params_json 列（跨店发布的目标店参数快照）。幂等。

    ⚠️ 为什么不写进 _migrate_add_shop_id：那个函数在「表里已有 shop_id」时直接 return，
       而所有现存库都早已是多店版，新列塞进去会永远补不上（老库启动即报 no such column）。
    ⚠️ 已有任务行回填 '{}'：表示「与批次同店发布」，即改造前的语义，行为完全不变。
    """
    _safe_add_column(conn, "publish_jobs", "params_json TEXT NOT NULL DEFAULT '{}'")


def _meta_get(conn, key, default=None):
    try:
        row = conn.execute("SELECT value FROM _meta WHERE key=?", (key,)).fetchone()
    except sqlite3.OperationalError:
        return default
    return row["value"] if row else default


def _meta_set(conn, key, value):
    conn.execute("INSERT OR REPLACE INTO _meta(key,value) VALUES(?,?)", (key, value))


def _reassign_legacy_shop(conn):
    """老库回填用的 shop_id 后来在 shops.json 里改了名时，自动把历史数据迁到新默认店。

    场景：运营先启动了新版本（此时还没有 shops.json，历史数据回填成 "default"），
    之后才按模板创建 shops.json（默认店叫 diamond_01）。若不处理，276 个批次会变成孤儿。

    只在安全条件下自动迁移：
      - 老 shop_id 已不再出现在 shops.json 里；
      - 目标（默认店）名下一行数据都没有（避免两个店的数据混在一起）。
    不满足条件时只打印警告，由人工确认，绝不擅自合并数据。
    """
    legacy = _meta_get(conn, "legacy_shop_id")
    if not legacy:
        return
    try:
        target = shop_registry.default_shop()["shop_id"]
    except ValueError:
        return
    if legacy == target:
        return
    if shop_registry.get_shop(legacy):
        return  # 老店仍在配置里，说明命名没变，不动

    tables = ("batches", "category_aliases", "publish_jobs", "publish_items", "image_upload_cache")
    occupied = conn.execute("SELECT COUNT(*) FROM batches WHERE shop_id=?", (target,)).fetchone()[0]
    if occupied:
        print(f"[db-migrate] WARN: 历史数据在 '{legacy}' 下，但默认店 '{target}' 已有数据，"
              f"未自动迁移。请人工确认后用 SQL 手动改名（5 张表的 shop_id 列）。")
        return

    moved = 0
    for table in tables:
        cur = conn.execute(f"UPDATE {table} SET shop_id=? WHERE shop_id=?", (target, legacy))
        moved += cur.rowcount or 0
    _meta_set(conn, "legacy_shop_id", target)
    print(f"[db-migrate] 历史数据已从 '{legacy}' 迁移到默认店 '{target}'（共 {moved} 行）")


def _json_post(url, payload, timeout=180, shop_id=None, extra_headers=None):
    headers = {"Content-Type": "application/json"}
    if shop_id:
        headers[SHOP_HEADER] = shop_id
    if extra_headers:
        headers.update(extra_headers)
    request = urllib.request.Request(url, data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                                     headers=headers, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        try:
            detail = json.loads(exc.read().decode("utf-8")).get("detail")
        except Exception:
            detail = None
        raise RuntimeError(str(detail or exc)[:1000]) from exc
    except Exception as exc:
        raise RuntimeError(str(exc)[:500]) from exc
    return data


def fit_wechat_title(title):
    """微信标题上限 60 个字符（官方文档：中文/字母/数字各算 1 个字符）。"""
    return str(title or "")[:60]


def fit_xhs_title(title):
    """小红书标题压到合法长度。

    平台报错原文：标题长度需要在[8]-[30]个字或[16]-[60]个字符
      「字」= 每个字符算 1；「字符」= 中文/全角算 2、字母数字算 1；满足任一区间即可。
      取「≤30 个字」这一档最稳（≤30 字时加权必然 ≤60 字符）。
    """
    text = str(title or "")
    if len(text) > 30:
        return text[:30]
    if len(text) < 8:
        return (text + " 培育钻石")[:30]
    return text


def platform_title(product, mapping, platform):
    """取要发给平台的标题，并做长度兜底。

    优先顺序：本平台标题列（运营在审核抽屉里改过的）→ 商品行里本平台的标题 → 通用标题。
    踩过的坑：只写 mapping.get("xhs_title") or product.get("title")，映射里没有小红书标题时
    就回落到「标题」= 微信那版（带品牌/系列，54 字/75 字符），触发 -5000500 标题长度错误，
    实测一个批次 250 个商品全因此失败。
    """
    column = "wechat_title" if platform == "wechat" else "xhs_title"
    raw = mapping.get(column) or product.get(column) or product.get("title", "")
    fixed = fit_wechat_title(raw) if platform == "wechat" else fit_xhs_title(raw)
    if fixed != str(raw or ""):
        print(f"[{platform}] 标题长度不合规，已裁剪为 {fixed!r}（原 {str(raw)!r}）")
    return fixed


def _wechat_payload(product, mapping, shop_id=None):
    chain = _latest_wechat_chain(mapping.get("wechat_category_chain") or [], shop_id)
    if not chain:
        raise RuntimeError("未配置微信类目")
    attrs = mapping.get("wechat_attrs") or product.get("attributes") or {}
    attr_list = [{"attr_key": str(k), "attr_value": str(v)} for k, v in attrs.items() if v not in (None, "", [])]
    freight_id = mapping.get("wechat_freight_template_id") or mapping.get("freight_template_id")
    if not freight_id:
        raise RuntimeError("未配置微信运费模板")
    if not product.get("main_images"):
        raise RuntimeError("微信商品至少需要一张主图")
    item = {
        "title": platform_title(product, mapping, "wechat"),
        "out_product_id": product.get("product_code"),
        "head_imgs": product.get("main_images") or [],
        "cats_v2": [{"cat_id": n.get("cat_id") or n.get("id")} for n in chain],
        "brand_id": str(mapping.get("wechat_brand_id") or "2100000000"),
        "express_info": {"template_id": str(freight_id)},
        "attrs": attr_list,
        "deliver_method": 0,
        "extra_service": {"seven_day_return": 1, "freight_insurance": 0},
    }
    if product.get("description") or product.get("detail_images"):
        item["desc_info"] = {}
        if product.get("description"): item["desc_info"]["detail"] = product["description"]
        if product.get("detail_images"): item["desc_info"]["imgs"] = product["detail_images"]
    skus = []
    for sku in product.get("skus", []):
        attrs_sku = [{"attr_key": s.get("name", ""), "attr_value": s.get("value", "")} for s in sku.get("specs", []) if s.get("name") and s.get("value")]
        item_sku = {"out_sku_id": sku.get("sku_code"), "sale_price": round(float(sku.get("price") or 0) * 100),
                    "market_price": round(float(sku.get("original_price") or sku.get("price") or 0) * 100),
                    "stock_num": int(float(sku.get("stock") or 0)), "sku_attrs": attrs_sku}
        # SKU 缩略图: 没匹配到图就不传, 避免空串被判非法
        if sku.get("uploaded_sku_image"):
            item_sku["thumb_img"] = sku["uploaded_sku_image"]
        skus.append(item_sku)
    item["skus"] = skus
    return item


def _latest_wechat_chain(saved_chain, shop_id=None):
    """用已存映射的叶子类目名, 回微信服务取最新 cats_v2 完整链路。

    ⚠️ 必须带 shop_id: 微信服务的类目树/凭证按店铺隔离, 不带头会落到默认店,
    发出品用的链路与该店实际权限可能对不上(默认店凭证异常时还会直接 400)。
    """
    if not saved_chain:
        return []
    leaf = saved_chain[-1]
    leaf_id = str(leaf.get("cat_id") or leaf.get("id") or "")
    leaf_name = str(leaf.get("name") or "").strip()
    if not leaf_name:
        raise RuntimeError("微信类目映射缺少叶子类目名称，请重新确认类目")
    cache_key = (shop_id, leaf_id, leaf_name)
    if cache_key in _WECHAT_CATEGORY_CHAIN_CACHE:
        return _WECHAT_CATEGORY_CHAIN_CACHE[cache_key]
    try:
        data = _http_get_json(f"{WECHAT_API_BASE}/categories?keyword={urllib.parse.quote(leaf_name)}", timeout=180, shop_id=shop_id)
    except Exception as exc:
        # 原来只带一句 "HTTP Error 400: Bad Request", 既看不出是哪个店也看不出原因;
        # 两个平台服务都是 FastAPI, 真实原因在响应体的 detail 里。
        detail = _http_error_detail(exc) if isinstance(exc, urllib.error.HTTPError) else str(exc)
        raise RuntimeError(f"无法校验最新微信 cats_v2 类目（店铺 {shop_id or '默认店'}）：{detail}") from exc
    candidates = [item for item in data.get("results", []) if item.get("leaf")]
    matched = [item for item in candidates if str(item.get("cat_id")) == leaf_id]
    if len(matched) != 1:
        saved_names = [str(node.get("name") or "").strip() for node in saved_chain]
        matched = [item for item in candidates if [str(node.get("name") or "").strip() for node in item.get("chain", [])] == saved_names]
    if len(matched) != 1:
        raise RuntimeError(f"微信叶子类目 {leaf_name} 无法在最新 cats_v2 中唯一确认，请重新选择类目")
    chain = matched[0].get("chain") or []
    _WECHAT_CATEGORY_CHAIN_CACHE[cache_key] = chain
    return chain


def _local_image(source):
    """Resolve an importer-owned local:// reference without allowing path traversal."""
    if not source.startswith("local://"):
        return None
    relative = source[len("local://"):].replace("\\", "/")
    parts = relative.split("/")
    if len(parts) != 2 or not re.fullmatch(r"[0-9a-f]{32}", parts[0]) or not re.fullmatch(r"[0-9a-f]{64}\.[a-z0-9]{2,5}", parts[1]):
        raise RuntimeError("Excel 图片引用无效")
    path = os.path.abspath(os.path.join(IMAGE_DIR, parts[0], parts[1]))
    root = os.path.abspath(IMAGE_DIR)
    if os.path.commonpath((root, path)) != root or not os.path.isfile(path):
        raise RuntimeError("Excel 内嵌图片文件不存在")
    with open(path, "rb") as image_file:
        content = image_file.read()
    if not content:
        raise RuntimeError("Excel 内嵌图片内容为空")
    return path, content


def _platform_image(platform, source_url, shop_id=None):
    global _image_cache_hit_mem, _image_cache_hit_db, _image_cache_miss
    # shop_id 缺失时回退默认店：新库 image_upload_cache.shop_id 是 NOT NULL，
    # 不兜底的话写入会直接失败。正常调用方(worker)始终显式传入。
    if not shop_id:
        try:
            shop_id = shop_registry.default_shop()["shop_id"]
        except ValueError:
            shop_id = "default"  # 极端情况：完全没有可用店铺时也要能写入
    source_url = str(source_url or "").strip()
    if not source_url: return ""
    local_image = _local_image(source_url)
    if not local_image and is_share_image(source_url):
        # 共享盘图片懒拷贝: 只有真正发布到平台时才读字节(扫描阶段零磁盘)
        local_image = fetch_share_image(source_url, shop_id)
    source_hash = hashlib.sha256(local_image[1] if local_image else source_url.encode("utf-8")).hexdigest()
    cache_key = (shop_id or "", platform, source_hash)
    # 1. 内存 LRU 层
    if cache_key in _IMAGE_CACHE_MEM:
        _image_cache_hit_mem += 1
        # 移到末尾表示最近使用
        _IMAGE_CACHE_ORDER.remove(cache_key)
        _IMAGE_CACHE_ORDER.append(cache_key)
        return _IMAGE_CACHE_MEM[cache_key]
    # 2. SQLite 持久层
    conn = db()
    cached = conn.execute("SELECT platform_url FROM image_upload_cache WHERE shop_id=? AND platform=? AND source_hash=?", (shop_id, platform, source_hash)).fetchone()
    conn.close()
    if cached:
        _image_cache_hit_db += 1
        _image_cache_mem_put(cache_key, cached["platform_url"])
        return cached["platform_url"]
    # 3. 未命中: 真正上传
    _image_cache_miss += 1
    if local_image:
        payload = {"filename": os.path.basename(local_image[0]),
                   "content_base64": base64.b64encode(local_image[1]).decode("ascii")}
        endpoint = "/images/upload-file" if platform == "wechat" else "/materials/upload-file"
        base = WECHAT_API_BASE if platform == "wechat" else XHS_API_BASE
        response = _json_post(f"{base}{endpoint}", payload, shop_id=shop_id)
    elif platform == "wechat":
        response = _json_post(f"{WECHAT_API_BASE}/images/upload", {"img_url": source_url}, shop_id=shop_id)
    else:
        response = _json_post(f"{XHS_API_BASE}/materials/upload", {"url": source_url}, shop_id=shop_id)
    if response.get("ok") is False: raise RuntimeError(str(response.get("detail") or response)[:500])
    result = response.get("result")
    if isinstance(result, str): platform_url = result
    else: platform_url = (result or {}).get("url") or (result or {}).get("materialUrl") or (result or {}).get("fileUrl") or (result or {}).get("img_url")
    if not platform_url: raise RuntimeError(f"{platform} 图片上传未返回地址")
    conn = db()
    conn.execute("INSERT OR REPLACE INTO image_upload_cache(shop_id,platform,source_hash,source_url,platform_url,updated_at) VALUES(?,?,?,?,?,?)",
                 (shop_id, platform, source_hash, source_url, platform_url, now()))
    conn.commit(); conn.close()
    _image_cache_mem_put(cache_key, platform_url)
    return platform_url


def _image_cache_mem_put(cache_key, platform_url):
    """把一条缓存写入内存层,超出上限时淘汰最久未用的。"""
    if cache_key in _IMAGE_CACHE_MEM:
        _IMAGE_CACHE_ORDER.remove(cache_key)
    _IMAGE_CACHE_MEM[cache_key] = platform_url
    _IMAGE_CACHE_ORDER.append(cache_key)
    while len(_IMAGE_CACHE_ORDER) > _IMAGE_CACHE_MAX:
        evicted = _IMAGE_CACHE_ORDER.pop(0)
        _IMAGE_CACHE_MEM.pop(evicted, None)


def _upload_images_parallel(platform: str, urls: list, max_workers: int = 10, shop_id=None) -> list:
    """并发上传图片并保序返回(替代原来的串行列表推导)。

    - 命中 image_upload_cache 的图片直接返回,不产生网络请求
    - 单张失败时抛出原始异常(与串行语义一致),发布项会标记 failed、可重试
    - max_workers 默认 10(平台限流风险可控,且缓存命中后实际请求远少于总数)
    """
    from concurrent.futures import ThreadPoolExecutor
    targets = [str(u or "").strip() for u in (urls or [])]
    targets = [u for u in targets if u]
    if not targets:
        return []
    if len(targets) == 1:
        return [_platform_image(platform, targets[0], shop_id=shop_id)]
    with ThreadPoolExecutor(max_workers=min(max_workers, len(targets))) as pool:
        # map 保序;任一任务抛异常会在迭代到该位置时重新抛出
        return list(pool.map(lambda u: _platform_image(platform, u, shop_id=shop_id), targets))


def _upload_image_uncached(platform, source_url, shop_id=None):
    """绕开 image_upload_cache 直接上传, 拿一个全新的平台素材地址。

    背景(实测): 小红书同一张图重复上传会返回**不同的**素材 URL, 而同一商品内多个 SKU
    共用同一个素材地址时, 小红书后台的「规格图」列显示为空; 图各不相同的商品则正常显示。
    所以只有"撞图"的 SKU 需要走这里补传一次, 没撞图的继续走缓存, 不做无谓重复上传。
    """
    local_image = _local_image(source_url)
    if not local_image and is_share_image(source_url):
        local_image = fetch_share_image(source_url, shop_id)
    if not local_image:
        raise RuntimeError(f"无法读取图片: {str(source_url)[:120]}")
    payload = {"filename": os.path.basename(local_image[0]),
               "content_base64": base64.b64encode(local_image[1]).decode("ascii")}
    endpoint = "/images/upload-file" if platform == "wechat" else "/materials/upload-file"
    base = WECHAT_API_BASE if platform == "wechat" else XHS_API_BASE
    response = _json_post(f"{base}{endpoint}", payload, shop_id=shop_id)
    if response.get("ok") is False:
        raise RuntimeError(str(response.get("detail") or response)[:300])
    result = response.get("result")
    if isinstance(result, str):
        return result
    return ((result or {}).get("url") or (result or {}).get("materialUrl")
            or (result or {}).get("fileUrl") or (result or {}).get("img_url") or "")


def _upload_sku_images(platform, skus, shop_id=None):
    """并发上传各 SKU 的规格图, 保序返回并保留空位。

    ⚠️ 不能直接用 _upload_images_parallel: 它会过滤空值导致返回列表变短,
       与 skus 按下标 zip 会串图(把 A 的图挂到 B 上)。
    ⚠️ 同一商品内多个 SKU 共用同一张图 → 同一个素材地址 → 小红书后台规格图列显示为空。
       所以发现撞图时, 给后面的 SKU 重新上传一次, 保证每个 SKU 地址唯一。
       (货盘图库每种颜色常只存一套图, 同色不同分数必然撞图, 这里就是处理这种情况。)
    """
    references = [str(sku.get("sku_image") or "").strip() for sku in skus]
    pending = [r for r in references if r]
    if not pending:
        return ["" for _ in references]
    queue = list(_upload_images_parallel(platform, pending, shop_id=shop_id))
    urls = [queue.pop(0) if r else "" for r in references]
    seen = {}
    for index, url in enumerate(urls):
        if not url:
            continue
        if url in seen:
            try:
                urls[index] = _upload_image_uncached(platform, references[index], shop_id=shop_id)
            except Exception as exc:
                print(f"[sku_image] 撞图重新上传失败, 保留原地址: {exc}")
        else:
            seen[url] = index
    return urls


def _prepare_platform_images(product, platform, shop_id=None):
    prepared = dict(product)
    prepared["main_images"] = _upload_images_parallel(platform, product.get("main_images", []), shop_id=shop_id)
    prepared["detail_images"] = _upload_images_parallel(platform, product.get("detail_images", []), shop_id=shop_id)
    # SKU 规格图: 微信 → skus[].thumb_img, 小红书 → sku_list[].specImage
    skus = [dict(sku) for sku in product.get("skus", [])]
    for sku, url in zip(skus, _upload_sku_images(platform, skus, shop_id=shop_id)):
        if url:
            sku["uploaded_sku_image"] = url
    # 小红书开了「主规格图」后要求每个规格值都有图(否则报「启用规格大图需上传所有规格图」)。
    # 图库缺图的 SKU(如某颜色没拍图)用商品主图兜底, 保证每个颜色值都有图。
    fallback = (prepared.get("main_images") or [""])[0]
    if platform == "xhs" and fallback:
        for sku in skus:
            if not sku.get("uploaded_sku_image"):
                sku["uploaded_sku_image"] = fallback
    prepared["skus"] = skus
    return prepared


def _get_xhs_var_candidates(category_id, var_id, shop_id=None):
    """拉小红书某个规格维度(如「颜色分类」)的候选值,带缓存(12h)。
    用于 SKU 规格值映射:货盘英文编码 → 平台中文 valueName + valueId。"""
    key = (shop_id, str(category_id), str(var_id))
    hit = _XHS_VAR_CANDIDATES_CACHE.get(key)
    if hit and time.time() - hit["ts"] < 12 * 3600:
        return hit["data"]
    try:
        url = (f"{XHS_API_BASE}/attribute-values"
               f"?category_id={urllib.parse.quote(str(category_id))}"
               f"&attribute_id={urllib.parse.quote(str(var_id))}")
        data = _http_get_json(url, timeout=60, shop_id=shop_id)
        result = data.get("result") or {}
        cands = result.get("attributeValueV3s") or result.get("values") or []
        _XHS_VAR_CANDIDATES_CACHE[key] = {"ts": time.time(), "data": cands}
        return cands
    except Exception:
        return []


def _resolve_xhs_spec_value(source_value, var_id, category_id, shop_id=None):
    """把货盘 SKU 规格值(可能英文)映射成小红书平台的中文 valueName + valueId。
    顺序:① 精确匹配候选值 ② 模糊匹配 ③ 别名表(英文→中文) ④ 原值兜底(不带 valueId)。"""
    sv = str(source_value).strip()
    cands = _get_xhs_var_candidates(category_id, var_id, shop_id) or []
    hit = next((c for c in cands if c.get("valueName") == sv), None)
    if not hit:
        hit = next((c for c in cands if sv in (c.get("valueName") or '') or (c.get("valueName") or '') in sv), None)
    if hit:
        return hit.get("valueName", sv), hit.get("valueId", '')
    # 别名兜底:货盘英文 SKU 编码 → 平台中文 valueName
    for cn, aliases in XHS_SPEC_VALUE_ALIASES.items():
        if sv in aliases:
            vhit = next((c for c in cands if c.get("valueName") == cn), None)
            if vhit:
                return cn, vhit.get("valueId", '')
            return cn, ''
    return sv, ''


def _xhs_payload(product, mapping, group_config=None, shop_id=None):
    category_id = mapping.get("xhs_category_id")
    if not category_id:
        raise RuntimeError("未配置小红书类目")
    if not mapping.get("xhs_brand_id"):
        raise RuntimeError("小红书品牌未唯一匹配")
    if not mapping.get("xhs_shipping_template_id"):
        raise RuntimeError("未选择小红书运费模板")
    if not mapping.get("xhs_logistics_plan_id"):
        raise RuntimeError("未选择小红书物流方案")
    if not product.get("main_images"):
        raise RuntimeError("小红书商品至少需要一张主图")
    item = {"name": platform_title(product, mapping, "xhs"),
            "brandId": str(mapping.get("xhs_brand_id") or ""), "categoryId": str(category_id),
            "attributes": [], "shippingTemplateId": str(mapping.get("xhs_shipping_template_id") or ""),
            "shippingGrossWeight": int(float(product.get("weight") or 500)),
            "variantIds": [], "images": product.get("main_images") or [],
            "imageDescriptions": product.get("detail_images") or [],
            "description": product.get("description", ""), "deliveryMode": "0", "freeReturn": "1",
            "articleNo": (product.get("product_code") or "").split("-", 1)[0]}
    attrs = mapping.get("xhs_attrs") or {}
    for key, val in attrs.items():
        values = val if isinstance(val, list) else [val]
        for value_item in values:
            if isinstance(value_item, dict):
                item["attributes"].append(value_item)
            elif value_item not in (None, ""):
                item["attributes"].append({"propertyId": str(key), "valueId": str(value_item), "value": str(value_item)})
    xhs_config = (group_config or {}).get("xhs") or {}
    var_defs = xhs_config.get("var_defs") or []
    spec_map = xhs_config.get("spec_map") or {}
    candidates = xhs_config.get("candidates") or {}
    mapped_var_ids = [str(v.get("id")) for v in var_defs if spec_map.get(str(v.get("id"))) or spec_map.get(v.get("id"))]
    item["variantIds"] = mapped_var_ids
    # 「规格大图」开关:平台要求开启该开关的商品必须至少有一个主规格项(variantIds 非空),
    # 否则 createItemV2 直接报 -5000500「当前商品未包含主规格项,无法开启规格大图功能」。
    # 因此只在确实映射到规格维度时才带;没有维度(如 spec_map 为空)就不带,商品照常发布(只是没有规格图列)。
    # 注:该字段只能在创建时带,审核中的商品 updateItemV2 改不动它。
    if mapped_var_ids:
        item["enableMainSpecImage"] = True
    sku_list = []
    for sku in product.get("skus", []):
        variants = []
        specs_by_name = {s.get("name"): s.get("value") for s in sku.get("specs", [])}
        for var_def in var_defs:
            var_id = str(var_def.get("id"))
            source_name = spec_map.get(var_id) or spec_map.get(var_def.get("id"))
            source_value = specs_by_name.get(source_name)
            if not source_name or not source_value: continue
            # SKU 规格值:货盘英文编码 → 平台中文 valueName + valueId(小红书后台才能显示中文)
            value_name, value_id = _resolve_xhs_spec_value(source_value, var_id, category_id, shop_id)
            variant = {"id": var_def.get("id"), "name": var_def.get("name"), "value": value_name}
            if value_id:
                variant["valueId"] = value_id
            variants.append(variant)
        item_sku = {"ipq": 1, "originalPrice": round(float(sku.get("original_price") or sku.get("price") or 0) * 100),
                    "price": round(float(sku.get("price") or 0) * 100), "stock": int(float(sku.get("stock") or 0)),
                    "logisticsPlanId": str(mapping.get("xhs_logistics_plan_id") or ""), "variants": variants,
                    "deliveryTime": {"time": "24", "type": "RELATIVE_TIME_NEW"}, "erpCode": sku.get("sku_code")}
        # SKU 规格图: 没匹配到图就不传
        if sku.get("uploaded_sku_image"):
            item_sku["specImage"] = sku["uploaded_sku_image"]
        sku_list.append(item_sku)
    return {"item": item, "sku_list": sku_list}


# 跨店发布：任务(job)里存了目标店的店铺私有参数，执行时必须覆盖批次里的来源店参数，
# 否则会把来源店的运费模板/物流方案/品牌发给目标店（平台会拒：模板不属于该店）。
# 只覆盖这 4 个「按店不同」的键；类目/属性/SKU规格是平台级的，跨店通用，绝不能动。
_SHOP_SCOPED_MAPPING_KEYS = (
    "wechat_freight_template_id",   # 微信运费模板
    "xhs_shipping_template_id",     # 小红书运费模板
    "xhs_logistics_plan_id",        # 小红书物流方案
    "xhs_brand_id",                 # 小红书品牌
)


def _job_params_of(row):
    """取任务所属 job 的店铺参数快照。

    row 由 worker 的 LEFT JOIN 查询产出；老库或同店发布时该字段可能是 NULL/空串，一律当空处理。
    """
    try:
        raw = row["job_params_json"]
    except (IndexError, KeyError):
        return {}
    if not raw:
        return {}
    try:
        data = json.loads(raw)
    except (TypeError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


# 发布前必须齐的「店铺级参数」: 缺了平台会整批拒。
# 以前是建完任务、worker 逐个商品报"未配置微信运费模板", 运营看到满屏失败、回头补选还得重发一轮;
# 现在建任务前直接拦住(见 publish 接口), 一个任务项都不建。
# 每项 = (该平台映射里可能用的键..., 给人看的名字)
_REQUIRED_SHOP_PARAMS = {
    "wechat": (("wechat_freight_template_id", "freight_template_id", "微信运费模板"),),
    "xhs": (("xhs_shipping_template_id", "小红书运费模板"), ("xhs_logistics_plan_id", "小红书物流方案")),
}


def _missing_shop_params(mappings, products, platforms, job_params=None):
    """列出这批商品发布前还缺哪些店铺级必填参数(空列表=齐全)。

    与 _wechat_payload/_xhs_payload 的必填判断保持同一套键: 跨店发布时先套用目标店参数快照,
    和 worker 真正发品时的取值口径一致, 避免"这里放行、那边照样失败"。
    """
    product_mappings = (mappings or {}).get("products", {}) or {}
    problems = []
    for platform in platforms:
        counts = {}
        for product in products:
            mapping = _apply_job_params(product_mappings.get(product.get("product_code")) or {}, job_params)
            for keys in _REQUIRED_SHOP_PARAMS.get(platform, ()):
                if not any(str(mapping.get(key) or "").strip() for key in keys[:-1]):
                    counts[keys[-1]] = counts.get(keys[-1], 0) + 1
        problems.extend(f"{label}（{count} 件商品未选）" for label, count in counts.items())
    return problems


def _apply_job_params(mapping, job_params):
    """把目标店的店铺私有参数覆盖到映射上。

    ⚠️ 必须返回副本：mapping 取自批次缓存(get_batch_cached)，就地改写会污染缓存，
       导致之后同一批次的「同店发布」也用上目标店的模板。
    """
    if not job_params:
        return mapping
    merged = dict(mapping)
    for key in _SHOP_SCOPED_MAPPING_KEYS:
        value = job_params.get(key)
        if value:
            merged[key] = value
    return merged


def _run_publish_item(row):
    shop_id = row["shop_id"]
    # 用缓存版:避免每个发布项都重新 json.loads 全量 rows + mappings(800 商品时单次近百毫秒)
    batch = get_batch_cached(row["batch_id"])
    product = next((p for p in batch["products"] if p.get("product_code") == row["product_code"]), None)
    if not product: raise RuntimeError("商品不存在")
    # 批次映射(来源店) + 任务参数(目标店) → 真正用于发布的映射
    mapping = _apply_job_params(
        batch["mappings"].get("products", {}).get(row["product_code"], {}),
        _job_params_of(row),
    )
    group_config = batch["mappings"].get("attr_groups", {}).get(product.get("internal_category"), {})
    if row["platform"] == "wechat":
        product = _prepare_platform_images(product, "wechat", shop_id=shop_id)
        result = _json_post(f"{WECHAT_API_BASE}/products", _wechat_payload(product, mapping, shop_id), shop_id=shop_id)
        if result.get("ok") is False: raise RuntimeError(str(result.get("detail") or result)[:500])
        value = result.get("result")
        return {"product_id": value} if not isinstance(value, dict) else value
    if row["platform"] == "xhs":
        product = _prepare_platform_images(product, "xhs", shop_id=shop_id)
        result = _json_post(f"{XHS_API_BASE}/items/and-sku", _xhs_payload(product, mapping, group_config, shop_id), shop_id=shop_id)
        value = result.get("result") or {}
        if value.get("itemId") and (result.get("partial") or value.get("skuErrors")):
            value["partial_error"] = json.dumps(value.get("skuErrors") or [], ensure_ascii=False)
            return value
        if result.get("ok") is False: raise RuntimeError(str(result.get("detail") or result)[:500])
        return value
    raise RuntimeError(f"不支持的平台: {row['platform']}")


def recover_interrupted_items(shop_id=None):
    """启动时把中断在 running 状态的项改回 queued。

    多店铺版：可选按 shop_id 恢复；不传则恢复所有店的中断项。
    """
    conn = db()
    if shop_id:
        conn.execute("UPDATE publish_items SET status='queued' WHERE status='running' AND shop_id=?", (shop_id,))
    else:
        conn.execute("UPDATE publish_items SET status='queued' WHERE status='running'")
    conn.commit(); conn.close()


def _platform_ready(platform, shop_id=None):
    base_url = WECHAT_API_BASE if platform == "wechat" else XHS_API_BASE
    try:
        url = f"{base_url}/health"
        if shop_id:
            req = urllib.request.Request(url, headers={SHOP_HEADER: shop_id})
        else:
            req = url
        with urllib.request.urlopen(req, timeout=3) as response:
            return response.status == 200
    except Exception:
        return False


def _worker_loop(platform=None, shop_id=None):
    """发布 worker 主循环。

    多店铺模式：
      - shop_id 非空：只处理该店铺的任务（按平台过滤，platform=None 则全平台）。
      - shop_id 为空：兼容旧行为，取默认店铺。
    """
    if not shop_id:
        try:
            shop_id = shop_registry.default_shop()["shop_id"]
        except ValueError:
            shop_id = "default"
    while True:
        conn = None
        try:
            if platform and not _platform_ready(platform, shop_id=shop_id):
                time.sleep(5)
                continue
            conn = db()
            conn.execute("BEGIN IMMEDIATE")
            # priority DESC: 分批发布(只选中几件)的项插到前面先跑, 再按 id 先到先服务
            # JOIN 出 job 的店铺参数快照：跨店发布时用它覆盖批次里的来源店参数
            if platform:
                item = conn.execute(
                    "SELECT i.*, j.params_json AS job_params_json FROM publish_items i "
                    "LEFT JOIN publish_jobs j ON j.id = i.job_id "
                    "WHERE i.shop_id=? AND i.status='queued' AND i.platform=? "
                    "ORDER BY i.priority DESC, i.id LIMIT 1",
                    (shop_id, platform)
                ).fetchone()
            else:
                item = conn.execute(
                    "SELECT i.*, j.params_json AS job_params_json FROM publish_items i "
                    "LEFT JOIN publish_jobs j ON j.id = i.job_id "
                    "WHERE i.shop_id=? AND i.status='queued' "
                    "ORDER BY i.priority DESC, i.id LIMIT 1",
                    (shop_id,)
                ).fetchone()
            if not item:
                conn.commit(); conn.close()
                time.sleep(1)
                continue
            conn.execute("UPDATE publish_items SET status='running', started_at=? WHERE id=? AND shop_id=?",
                         (now(), item["id"], shop_id)); conn.commit(); conn.close()
            try:
                result = _run_publish_item(item)
                product_id = result.get("product_id") or result.get("itemId") or result.get("id")
                sku_ids = result.get("skuIds") or []
                status = "partial" if result.get("partial_error") else "success"
                conn = db(); conn.execute(
                    "UPDATE publish_items SET status=?, platform_product_id=?, platform_sku_ids=?, error=?, finished_at=? "
                    "WHERE id=? AND shop_id=?",
                    (status, str(product_id or ""), json.dumps(sku_ids), result.get("partial_error"),
                     now(), item["id"], shop_id)); conn.commit(); conn.close()
            except Exception as exc:
                conn = db(); conn.execute(
                    "UPDATE publish_items SET status='failed', error=?, finished_at=? WHERE id=? AND shop_id=?",
                    (str(exc)[:1000], now(), item["id"], shop_id)); conn.commit(); conn.close()
            conn = db(); conn.execute("""UPDATE publish_jobs SET processed=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status IN ('success','failed','partial')), success=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status='success'), failed=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status IN ('failed','partial')), status=CASE WHEN (SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status IN ('queued','running'))=0 THEN CASE WHEN (SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status IN ('failed','partial'))>0 THEN 'completed_with_errors' ELSE 'completed' END ELSE 'running' END, updated_at=? WHERE id=? AND shop_id=?""",
                (shop_id, item["job_id"], shop_id, item["job_id"], shop_id, item["job_id"],
                 shop_id, item["job_id"], shop_id, item["job_id"], now(), item["job_id"], shop_id))
            conn.commit(); conn.close()
        except Exception:
            # 异常时必须回滚并关闭连接,否则 BEGIN IMMEDIATE 持有的写锁会一直不释放,锁死整个库
            if conn is not None:
                try:
                    conn.rollback()
                except Exception:
                    pass
                try:
                    conn.close()
                except Exception:
                    pass
            time.sleep(2)


class ImportBody(BaseModel):
    filename: str = "products.xlsx"
    content_base64: str
    zip_base64: str | None = None
    folder_files: list[dict] | None = None
    common_detail_files: list[dict] | None = None


class HuopaiImportBody(BaseModel):
    """货盘直读导入:读服务器本地货盘表,后台转换后入库(不走文件上传)。

    路径收敛: 只接受「货盘目录(HUOPAI_DIR,默认取 HUOPAI_PATH 所在目录)」内的 .xlsx,
    目录由服务端配置决定,不接受客户端指定任意目录。
    """
    path: str | None = None        # 货盘目录内的 .xlsx 路径;不传则用默认货盘表


class ImageScanBody(BaseModel):
    """图片直读:扫描图片根目录, 把商品图/SKU 图挂到批次上(不走浏览器上传, 几百个商品也没压力)。"""
    root: str | None = None            # 必须是服务端白名单内的图片根目录;不传用默认路径
    months: list[str] | None = None    # 只扫指定月份目录(如 ["7-9月份"]);None = 全部
    dry_run: bool = False              # True=只匹配不拷贝(先看命中率)


class ItemsBody(BaseModel):
    items: list[dict]


class SkuImageBody(BaseModel):
    """补规格图: 运营在批次表里给缺图的 SKU 单张上传(JSON+Base64, 与本项目其它上传接口一致)。"""
    content_base64: str
    filename: str | None = None


class MappingsBody(BaseModel):
    mappings: dict


class PublishBody(BaseModel):
    platforms: list[str]
    product_codes: list[str] | None = None  # 不传=全量发布,传了=只发布指定商品(分批发布)
    # 跨店发布：目标店铺（不传=当前店铺，即改造前的同店发布行为）。
    # 任务写到目标店名下(worker 用该店凭证发布)，params 是该店的私有参数
    # （微信运费模板 / 小红书运费模板+物流方案+品牌），执行时覆盖批次里的来源店参数。
    target_shop_id: str | None = None
    params: dict | None = None


class RetryBody(BaseModel):
    item_ids: list[int] | None = None


class AliasBody(BaseModel):
    internal_category: str
    wechat: dict | None = None
    xhs: dict | None = None


class AliasImportBody(BaseModel):
    filename: str = "category-aliases.xlsx"
    content_base64: str


def now():
    return datetime.now(timezone.utc).isoformat()


def value(row, *names):
    for name in names:
        if name in row and row[name] not in (None, ""):
            return str(row[name]).strip()
    return ""


def _embedded_image_extension(image, content):
    image_format = str(getattr(image, "format", "") or "").lower()
    if image_format in ("jpeg", "jpg") or content.startswith(b"\xff\xd8"):
        return "jpg"
    if image_format == "png" or content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if image_format == "gif" or content.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if image_format == "webp" or (content.startswith(b"RIFF") and content[8:12] == b"WEBP"):
        return "webp"
    return "png"


def _extract_embedded_images(sheet, headers, batch_id):
    column_types = {}
    for index, header in enumerate(headers):
        if header in ("主图", "主图链接"):
            column_types[index] = "main_images"
        elif header in ("详情图", "详情图链接"):
            column_types[index] = "detail_images"
    extracted = {}
    batch_dir = os.path.join(IMAGE_DIR, batch_id)
    for image in getattr(sheet, "_images", []):
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        if marker is None or marker.col not in column_types or marker.row < 1:
            continue
        try:
            content = image._data()
        except Exception as exc:
            raise HTTPException(400, f"Excel 第 {marker.row + 1} 行内嵌图片读取失败：{exc}")
        digest = hashlib.sha256(content).hexdigest()
        extension = _embedded_image_extension(image, content)
        os.makedirs(batch_dir, exist_ok=True)
        path = os.path.join(batch_dir, f"{digest}.{extension}")
        if not os.path.exists(path):
            with open(path, "wb") as image_file:
                image_file.write(content)
        reference = f"local://{batch_id}/{digest}.{extension}"
        values = extracted.setdefault(marker.row + 1, {}).setdefault(column_types[marker.col], [])
        if reference not in values:
            values.append(reference)
    return extracted


def parse_excel(raw: bytes, filename: str, batch_id=None):
    extension = os.path.splitext(filename or "")[1].lower()
    if extension == ".xls":
        raise HTTPException(400, "暂不支持旧版 .xls，请另存为 .xlsx 后再导入")
    if extension == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            return [normalize_row(row, index) for index, row in enumerate(reader, 2) if any(str(v or "").strip() for v in row.values())]
        except UnicodeDecodeError:
            raise HTTPException(400, "CSV 请保存为 UTF-8 编码后再导入")
        except Exception as exc:
            raise HTTPException(400, f"CSV 解析失败：{exc}")
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        embedded = _extract_embedded_images(sheet, headers, batch_id) if batch_id else {}
        result = []
        for index, cells in enumerate(rows[1:], 2):
            if not any(cell not in (None, "") for cell in cells):
                continue
            item = {headers[i]: cells[i] if i < len(cells) else "" for i in range(len(headers)) if headers[i]}
            normalized = normalize_row(item, index)
            for field, references in embedded.get(index, {}).items():
                normalized[field] = list(dict.fromkeys(normalized.get(field, []) + references))
            result.append(normalized)
        return result
    except Exception as exc:
        raise HTTPException(400, f"Excel 解析失败：{exc}")


def _parse_attributes(text):
    """解析商品属性列: "材质=18K金;钻石切工=Very good" → {"材质": "18K金", "钻石切工": "Very good"}"""
    attrs = {}
    if not text:
        return attrs
    for pair in re.split(r"[;；]", text):
        pair = pair.strip()
        if "=" in pair:
            k, v = pair.split("=", 1)
            if k.strip():
                attrs[k.strip()] = v.strip()
    return attrs


_CN_NUM = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "七", 8: "八", 9: "九", 10: "十"}


def _parse_specs(row, max_specs=20):
    """动态读取 规格N名称/规格N值 列（N 从 1 递增），返回 [{"name":.., "value":..}]。
    表头支持阿拉伯数字与中文数字（规格1名称 / 规格一名称）；只收集填了名称或值的规格列，按列序排列。
    规格维度数量不写死，运营可自由增减 Excel 规格列。"""
    specs = []
    for n in range(1, max_specs + 1):
        cn = _CN_NUM.get(n, str(n))
        name = value(row, f"规格{n}名称", f"规格{cn}名称")
        val = value(row, f"规格{n}值", f"规格{cn}值")
        if name or val:
            specs.append({"name": name, "value": val})
    return specs


def normalize_row(row, line):
    sku_code = value(row, "SKU编码", "SKU 编码", "sku_code", "erpCode")
    product_code = value(row, "商品编码", "商品 编码", "product_code", "articleNo")
    price = value(row, "售价", "售价（元）", "price", "price_yuan")
    stock = value(row, "库存", "stock")
    errors = []
    if not product_code:
        errors.append("缺少商品编码")
    if not value(row, "标题", "商品标题", "name"):
        errors.append("缺少标题")
    if not sku_code:
        errors.append("缺少SKU编码")
    try:
        if float(price) <= 0:
            errors.append("售价必须大于0")
    except (TypeError, ValueError):
        errors.append("售价格式错误")
    try:
        if int(float(stock)) < 0:
            errors.append("库存不能小于0")
    except (TypeError, ValueError):
        errors.append("库存格式错误")
    normalized = {
        "line": line, "product_code": product_code,
        "title": value(row, "标题", "商品标题", "name"),
        "wechat_title": value(row, "微信标题"), "xhs_title": value(row, "小红书标题"),
        "internal_category": value(row, "内部类目", "类目"),
        "brand": value(row, "品牌"), "description": value(row, "描述", "商品描述"),
        "attributes": _parse_attributes(value(row, "商品属性", "属性", "attributes")),
        "weight": value(row, "重量", "商品毛重"),
        "main_images": value(row, "主图", "主图链接").split(",") if value(row, "主图", "主图链接") else [],
        "detail_images": value(row, "详情图", "详情图链接").split(",") if value(row, "详情图", "详情图链接") else [],
        "sku_code": sku_code, "specs": _parse_specs(row),
        "original_price": value(row, "原价", "原价（元）"),
        "price": price, "stock": stock, "errors": errors,
    }
    # 货盘直读的来源定位(哪张表哪一行/原始商家编码)要原样带下去:
    # 发布完「回填平台商品ID到货盘表」靠它定位原文件的单元格。
    for key in ("_源表", "_源行", "_源码", "_系列原文"):
        if key in row:
            normalized[key] = row[key]
    return normalized


# 「图片还在共享盘/本地磁盘上」的引用前缀: 扫描阶段只记路径, 真正上传平台时才读字节落盘(懒拷贝)
SHARE_PREFIX = "disk:"


def is_share_image(source):
    """是否是共享盘图片引用(懒加载), 形如 'disk:\\\\192.168.10.250\\...\\a.jpg'。"""
    return str(source or "").startswith(SHARE_PREFIX)


def fetch_share_image(source, shop_id=None):
    """读共享盘上的图片并落盘缓存, 返回 (本地路径, 内容)。

    只在真正发布时调用: 扫描 233 个商品只记路径(秒级、零磁盘),
    发到哪个商品才读哪个商品的图, 不发布的商品完全不占磁盘。
    落盘按内容 sha256 命名, 同一张图(通用详情图/同款多 SKU 共用图)只存一份。

    安全: disk: 引用会存进批次数据, 而批次数据可由客户端通过 PUT /import/{id}/items
    覆写, 因此这里必须再校验一次路径白名单, 否则可被用来读服务器任意文件并上传平台。
    白名单必须按店铺取(见 _allowed_roots_for)，否则服务器上挂载点路径会被误拒。
    """
    path = str(source)[len(SHARE_PREFIX):]
    if not path or not _within_allowed_roots(path, _allowed_roots_for(shop_id)):
        raise RuntimeError(f"共享盘图片路径不在允许的图片目录内：{path}")
    # 用 os.stat 而不是 os.path.isfile: 后者把"无权限 / 网络不可达"也一并吞成 False,
    # 让权限问题伪装成"文件不存在"(线上踩过: worker 以 SYSTEM 跑时整批报"不存在")。
    try:
        os.stat(path)
    except FileNotFoundError:
        raise RuntimeError(f"共享盘图片不存在(可能已被移动或删除)：{path}")
    except PermissionError:
        raise RuntimeError(f"共享盘图片无权读取(Worker 运行账户对该共享没有权限)：{path}")
    except OSError as exc:
        raise RuntimeError(f"共享盘不可访问：{path}（{exc}）")
    if not os.path.isfile(path):
        raise RuntimeError(f"共享盘路径不是图片文件：{path}")
    with open(path, "rb") as image_file:
        content = image_file.read()
    if not content:
        raise RuntimeError(f"共享盘图片内容为空：{path}")
    digest = hashlib.sha256(content).hexdigest()
    extension = os.path.splitext(path)[1].lower() or ".jpg"
    target_dir = os.path.join(IMAGE_DIR, "_share")
    os.makedirs(target_dir, exist_ok=True)
    target = os.path.join(target_dir, f"{digest}{extension}")
    if not os.path.exists(target):
        with open(target, "wb") as image_file:
            image_file.write(content)
    return target, content


def _store_folder_images(files, batch_id, rows, common_detail_files=None):
    """Save browser-selected folder files and resolve Excel image filenames."""
    batch_dir = os.path.join(IMAGE_DIR, batch_id)
    os.makedirs(batch_dir, exist_ok=True)
    references = {}
    file_records = []
    common_details = []
    for item in [*(files or []), *(common_detail_files or [])]:
        relative = str(item.get("path") or item.get("name") or "").replace("\\", "/").strip()
        encoded = item.get("content_base64") or ""
        if not relative or not encoded:
            continue
        parts = [part for part in relative.split("/") if part not in ("", ".")]
        if any(part == ".." for part in parts):
            continue
        name = os.path.basename(relative)
        try:
            content = base64.b64decode(encoded, validate=True)
        except Exception:
            continue
        if not name or not content or len(content) > 20 * 1024 * 1024:
            continue
        digest = hashlib.sha256(content).hexdigest()
        extension = os.path.splitext(name)[1].lower() or ".jpg"
        target = os.path.join(batch_dir, f"{digest}{extension}")
        if not os.path.exists(target):
            with open(target, "wb") as image_file:
                image_file.write(content)
        reference = f"local://{batch_id}/{digest}{extension}"
        references[relative.lower()] = reference
        references.setdefault(name.lower(), reference)
        file_records.append((name, reference))
        if item in (common_detail_files or []):
            common_details.append(reference)
    for row in rows:
        for field in ("main_images", "detail_images"):
            row[field] = [
                references.get(str(source or "").strip().replace("\\", "/").lower(),
                               references.get(os.path.basename(str(source or "").strip()).lower(), source))
                for source in row.get(field, [])
            ]
        code = str(row.get("product_code") or "").strip().lower()
        # 取基础商品编码(去掉 -系列 后缀)用于匹配文件名,这样 KGN1000330主图1.jpg 能匹配 KGN1000330-永恒
        base_code = code.split("-", 1)[0] if code else ""
        if base_code:
            matches = [(name, ref) for name, ref in file_records if os.path.splitext(name)[0].lower().startswith(base_code)]
            if not row.get("main_images"):
                main_candidates = [ref for name, ref in sorted(matches) if "主图" in name or "main" in name.lower()]
                if not main_candidates and matches:
                    # 没有「主图/main」字样时,取匹配到的第一张当主图
                    main_candidates = [sorted(matches)[0][1]]
                row["main_images"] = main_candidates
            if not row.get("detail_images"):
                detail_candidates = [ref for name, ref in sorted(matches) if "详情" in name or "detail" in name.lower()]
                if not detail_candidates and len(matches) > 1:
                    # 没有「详情/detail」字样且多张时,除主图外都当详情图
                    detail_candidates = [ref for name, ref in sorted(matches)[1:]]
                row["detail_images"] = detail_candidates
            for reference in common_details:
                if reference not in row["detail_images"]:
                    row["detail_images"].append(reference)
            # 兜底:商品没匹配到主图时,优先用本次批次里 1:1 且 ≥800 的图(KGN 这类),其次才用通用详情
            # 1:1 + ≥800 才能过小红书 createItemV2 的「图片像素不低于800x800 (1:1)」
            if not row.get("main_images"):
                fallback = None
                # 先从文件夹图片里挑 1:1 且 ≥800 的
                for _n, _r in file_records:
                    try:
                        from PIL import Image as _Img
                        import io as _io
                        _p = _local_image(_r)
                        if not _p: continue
                        with _Img.open(_io.BytesIO(_p[1])) as _im:
                            _w, _h = _im.size
                            if _w >= 800 and _h == _w:
                                fallback = _r
                                break
                    except Exception:
                        continue
                # 再从通用详情里挑 1:1 且 ≥800 的
                if not fallback:
                    for _r in common_details:
                        try:
                            _p = _local_image(_r)
                            if not _p: continue
                            with _Img.open(_io.BytesIO(_p[1])) as _im:
                                _w, _h = _im.size
                                if _w >= 800 and _h == _w:
                                    fallback = _r
                                    break
                        except Exception:
                            continue
                # 实在没有 1:1 的,才用第一张通用详情(大概率会被小红书拒,生产应避免)
                if not fallback and common_details:
                    fallback = common_details[0]
                if fallback:
                    row["main_images"] = [fallback]


def group_products(rows):
    """将标准模板的 SKU 行合并为商品，供预览和后续平台发布使用。"""
    grouped = {}
    for row in rows:
        code = row.get("product_code") or f"__line_{row.get('line')}"
        product = grouped.setdefault(code, {
            "product_code": code, "title": row.get("title", ""),
            "wechat_title": row.get("wechat_title", ""), "xhs_title": row.get("xhs_title", ""),
            "internal_category": row.get("internal_category", ""), "brand": row.get("brand", ""),
            "description": row.get("description", ""), "attributes": row.get("attributes", {}),
            "weight": row.get("weight", ""),
            "main_images": row.get("main_images", []), "detail_images": row.get("detail_images", []),
            "skus": [], "errors": [], "warnings": [], "lines": [],
        })
        product["lines"].append(row.get("line"))
        product["errors"].extend(row.get("errors", []))
        for image_field in ("main_images", "detail_images"):
            for image_source in row.get(image_field, []):
                if image_source and image_source not in product[image_field]:
                    product[image_field].append(image_source)
        # 逐行留存商品属性，用于后续一致性检测（商品属性是 SPU 级，只采用第一行）
        product.setdefault("_attr_variants", []).append(row.get("attributes") or {})
        sku_key = row.get("sku_code")
        if any(sku.get("sku_code") == sku_key for sku in product["skus"]):
            product["errors"].append("同一商品内 SKU 编码重复")
        # 兼容旧批次 rows_json（早期为 spec1_name/spec1_value/spec2_*），统一转成 specs 数组
        row_specs = row.get("specs")
        if row_specs is None:
            row_specs = []
            for i in (1, 2):
                nm = row.get(f"spec{i}_name", "")
                vv = row.get(f"spec{i}_value", "")
                if nm or vv:
                    row_specs.append({"name": nm, "value": vv})
        product["skus"].append({
            "sku_code": sku_key,
            "specs": [{"name": sp.get("name", ""), "value": sp.get("value", "")} for sp in row_specs],
            "original_price": row.get("original_price", ""),
            "price": row.get("price", ""), "stock": row.get("stock", ""),
            # SKU 规格图(图片直读扫描写入, 固定为该 SKU 的「主图2」)
            "sku_image": row.get("sku_image", ""),
        })
    products = list(grouped.values())
    for product in products:
        product["errors"] = list(dict.fromkeys(product["errors"]))
        # 商品属性在同一商品的各 SKU 行间不一致时告警：
        # 平台模型只支持 SPU 级属性（微信 attrs / 小红书 attributes），
        # 实际仅第一行的值生效，其余行静默丢弃，必须提醒运营。
        variants = product.pop("_attr_variants", [])
        if len(variants) > 1:
            keys = set()
            for variant in variants:
                keys.update(variant.keys())
            for key in sorted(keys):
                values = [str(variant.get(key, "")).strip() for variant in variants]
                distinct = list(dict.fromkeys([v for v in values if v]))
                if len(distinct) > 1:
                    product["warnings"].append(
                        f"商品属性「{key}」在 SKU 间不一致（{' / '.join(distinct)}），"
                        f"当前仅采用第一行的值；若该差异影响选购或定价，建议提升为规格维度"
                    )
        product["warnings"] = list(dict.fromkeys(product.get("warnings", [])))
        # 汇总规格维度（按名称去重、保持列序），并把每个 SKU 的 specs 对齐到统一维度，
        # 避免各 SKU 规格列数量不同导致前端按下标渲染错位。
        dimensions = []
        for sku in product["skus"]:
            for sp in sku.get("specs", []):
                if sp.get("name") and sp["name"] not in dimensions:
                    dimensions.append(sp["name"])
        for sku in product["skus"]:
            by_name = {sp.get("name"): sp.get("value", "") for sp in sku.get("specs", []) if sp.get("name")}
            sku["specs"] = [{"name": d, "value": by_name.get(d, "")} for d in dimensions]
        product["spec_dimensions"] = dimensions
        if len(product["skus"]) > 100:
            product["errors"].append("单商品最多支持100个 SKU")
        titles = {rows[line - 2].get("title") for line in product["lines"] if isinstance(line, int) and 1 < line <= len(rows) + 1}
        if len(titles) > 1:
            product["errors"].append("同一商品标题不一致")
        prices = [float(s["price"]) for s in product["skus"] if str(s.get("price", "")).replace(".", "", 1).isdigit()]
        stocks = [int(float(s["stock"])) for s in product["skus"] if str(s.get("stock", "")).replace(".", "", 1).isdigit()]
        product["sku_count"] = len(product["skus"])
        product["price_min"] = min(prices) if prices else None
        product["price_max"] = max(prices) if prices else None
        product["stock_total"] = sum(stocks) if stocks else None
    return products


@app.get("/health")
def health():
    return {"status": "ok", "service": "bulk-catalog-api"}


@app.get("/shops")
def list_shops():
    """返回脱敏后的店铺列表，供前端店铺选择器使用。"""
    return {"ok": True, "result": shop_registry.redacted_list(only_enabled=True)}


@app.get("/template")
def template():
    headers = ["商品编码", "标题", "微信标题", "小红书标题", "内部类目", "品牌", "描述", "商品属性", "重量", "主图", "详情图", "SKU编码", "规格1名称", "规格1值", "规格2名称", "规格2值", "规格3名称", "规格3值", "原价（元）", "售价（元）", "库存"]
    try:
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "商品SKU"
        sheet.append(headers)
        sheet.append(["SPU-001", "示例商品标题", "", "", "珠宝", "", "", "材质=18K金;颜色=白色", 500, "", "", "SKU-001", "颜色", "红色", "尺码", "均码", "净度", "VS1", 99.9, 88.0, 100])
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        sheet.freeze_panes = "A2"
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(iter([output.read()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=bulk-product-template.xlsx"})
    except Exception as exc:
        raise HTTPException(500, f"生成 Excel 模板失败：{exc}")


@app.post("/import")
def import_batch(body: ImportBody, request: Request):
    shop_id = _current_shop_id(request)
    _verify_shop(shop_id)
    operator = _current_operator(request)
    try:
        raw = base64.b64decode(body.content_base64)
    except Exception:
        raise HTTPException(400, "Excel 文件编码无效")
    batch_id = uuid.uuid4().hex
    rows = parse_excel(raw, body.filename, batch_id)
    if len(rows) > 5000:
        raise HTTPException(400, "单批最多导入5000行SKU")
    if body.zip_base64:
        try:
            with zipfile.ZipFile(io.BytesIO(base64.b64decode(body.zip_base64))) as archive:
                for info in archive.infolist():
                    if info.is_dir() or info.filename.startswith(("/", "\\")) or ".." in info.filename:
                        continue
                    target = os.path.join(IMAGE_DIR, f"{uuid.uuid4().hex}_{os.path.basename(info.filename)}")
                    with open(target, "wb") as output:
                        output.write(archive.read(info))
        except Exception as exc:
            raise HTTPException(400, f"图片ZIP处理失败：{exc}")
    if body.folder_files or body.common_detail_files:
        _store_folder_images(body.folder_files, batch_id, rows, body.common_detail_files)
    errors = sum(bool(row["errors"]) for row in rows)
    conn = db()
    conn.execute("INSERT INTO batches(id,shop_id,filename,status,total,valid,errors,created_at,operator,rows_json) VALUES(?,?,?,?,?,?,?,?,?,?)",
                 (batch_id, shop_id, body.filename, "待校验" if errors else "待发布", len(rows), len(rows) - errors, errors, now(), operator, json.dumps(rows, ensure_ascii=False)))
    conn.commit(); conn.close()
    # 浏览器上传的 Excel 在服务器上没有原文件, 记一条空路径(回填ID到货盘表时会提示不支持)
    save_source_info(batch_id, "", body.filename)
    products = group_products(rows)
    product_errors = sum(bool(product["errors"]) for product in products)
    return {"ok": True, "batch_id": batch_id, "total": len(rows), "product_count": len(products), "valid": len(rows) - errors, "errors": errors, "product_errors": product_errors}


def _huopai_module():
    """货盘转换模块(局部导入:只有货盘相关接口需要 openpyxl 与转换规则)。"""
    try:
        import huopai_adapter as module
    except Exception as exc:
        raise HTTPException(500, f"货盘转换模块不可用：{exc}")
    return module


def _huopai_dir():
    """货盘表目录白名单根: HUOPAI_DIR 优先, 否则取默认货盘表所在目录。"""
    directory = os.environ.get("HUOPAI_DIR") or os.path.dirname(_huopai_module().HUOPAI_PATH)
    return os.path.abspath(directory)


def _resolve_huopai_file(path):
    """把请求里的货盘表路径收敛到白名单目录内的 .xlsx。

    防止 /import-huopai 被用来读取(并解析)服务器上任意文件: 之前 path 完全来自请求体,
    传 C:\\Windows\\... 也会被当成 Excel 打开。现在只允许货盘目录内的 .xlsx。
    """
    candidate = os.path.abspath(str(path or ""))
    root = _huopai_dir()
    if os.path.splitext(candidate)[1].lower() != ".xlsx":
        raise HTTPException(400, "货盘表只支持 .xlsx 文件")
    try:
        inside = os.path.commonpath((root, candidate)) == root
    except ValueError:
        inside = False
    if not inside:
        raise HTTPException(400, f"货盘表路径不在允许的货盘目录内：{root}")
    if not os.path.isfile(candidate):
        raise HTTPException(400, f"货盘表不存在：{candidate}")
    return candidate


@app.get("/huopai-files")
def list_huopai_files():
    """【货盘文件列表】列出货盘目录下可导入的 .xlsx,供前端下拉选择。

    目录由服务端配置(HUOPAI_DIR / HUOPAI_PATH)决定,接口不接受客户端传目录;
    因此运营侧的货盘表只要放进该目录即可,不再需要把绝对路径写进前端。
    """
    directory = _huopai_dir()
    if not os.path.isdir(directory):
        raise HTTPException(400, f"货盘目录不可访问：{directory}(请配置 HUOPAI_DIR / HUOPAI_PATH)")
    with os.scandir(directory) as scanner:
        entries = sorted(scanner, key=lambda entry: entry.name)
    files = []
    for entry in entries:
        if not entry.is_file() or os.path.splitext(entry.name)[1].lower() != ".xlsx":
            continue
        stat = entry.stat()
        files.append({"name": entry.name, "path": entry.path, "size": stat.st_size,
                      "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()})
    default_path = os.path.abspath(_huopai_module().HUOPAI_PATH)
    return {"ok": True, "dir": directory,
            "default": default_path if os.path.isfile(default_path) else "",
            "files": files}


@app.post("/import-huopai")
def import_huopai(body: HuopaiImportBody | None = None, request: Request = None):
    """【直读货盘】从服务器本地货盘目录读货盘表,后台跑「货盘→标准模板」转换后入库。
    与大平台一致:数据源直连,不经过浏览器上传,规避 64MB 大文件 Base64 传输。
    安全: 只允许读取货盘目录白名单内的 .xlsx(见 _resolve_huopai_file)。
    转换规则与 huopai_adapter.py 完全一致(编码/价格/库存/属性/图片/系列拆分)。"""
    shop_id = _current_shop_id(request) if request else shop_registry.default_shop()["shop_id"]
    _verify_shop(shop_id)
    operator = _current_operator(request) if request else shop_registry.default_operator()
    import openpyxl  # 局部导入:仅本接口需要
    H = _huopai_module()
    path = _resolve_huopai_file((body.path if body and body.path else None) or H.HUOPAI_PATH)
    # ① 转换:货盘 → 标准列 dict(复用 huopai_adapter 的业务规则)
    try:
        images = H.load_images(shop_id)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            raw_rows = []
            if "培育钻" in wb.sheetnames:
                raw_rows.extend(H.parse_peiyuzuan(wb["培育钻"], images))
            if "天然钻" in wb.sheetnames:
                raw_rows.extend(H.parse_tianranzuan(wb["天然钻"], images))
        finally:
            wb.close()
        H.enrich_carat(raw_rows)
        H.enrich_split_tone(raw_rows)
        H.enrich_split_chain(raw_rows)
        # 标题按商品编码生成, 必须在所有拆分之后(拆出来的商品标题要靠它区分开)
        H.enrich_titles(raw_rows)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(500, f"货盘转换失败：{exc}")
    if not raw_rows:
        raise HTTPException(400, "货盘表没解析到商品行(检查 sheet 名是否含「培育钻」「天然钻」)")
    # ② 复用标准导入的归一化 + 批次入库
    batch_id = uuid.uuid4().hex
    rows = [normalize_row(r, i) for i, r in enumerate(raw_rows, 2)
            if any(str(v or "").strip() for v in r.values())]
    if len(rows) > 5000:
        raise HTTPException(400, "单批最多导入5000行SKU")
    errors = sum(bool(row["errors"]) for row in rows)
    filename = os.path.basename(path)
    conn = db()
    conn.execute("INSERT INTO batches(id,shop_id,filename,status,total,valid,errors,created_at,operator,rows_json) VALUES(?,?,?,?,?,?,?,?,?,?)",
                 (batch_id, shop_id, filename, "待校验" if errors else "待发布", len(rows), len(rows) - errors, errors, now(), operator, json.dumps(rows, ensure_ascii=False)))
    conn.commit(); conn.close()
    # 记住货盘表原文件位置: 发布完「回填商品ID」直接改这个文件
    save_source_info(batch_id, path, filename)
    products = group_products(rows)
    return {"ok": True, "batch_id": batch_id, "filename": filename, "total": len(rows),
            "product_count": len(products), "valid": len(rows) - errors, "errors": errors,
            "product_errors": sum(bool(product["errors"]) for product in products)}


# ==================================================================
# 回填平台商品ID到货盘表原文件
# ==================================================================
# 发布完成后运营要把平台生成的商品ID登记回货盘表(培育钻的「微信小店」「小红书产品id」列),
# 原先只能照着导出的 CSV 一行行抄。这里直接改写货盘表原文件:
#   - 只改需要的单元格: 在 xlsx 内部做字节级替换, 不重新保存整本工作簿 ——
#     货盘表里 WPS 的内嵌图片(DISPIMG 主图)用 openpyxl 重存会丢;
#   - 写前自动备份原文件, 写完自检(重新读回单元格)通过才替换, 失败则原文件不动;
#   - 只填空白单元格: 已有值(人工填过)的行跳过并回报, 不覆盖;
#   - 逐行用导入时记下的商家编码核对, 货盘表被改过/重排过的行不写(宁可不写, 不写错行)。
SOURCE_INFO_DIR = os.path.join(BASE_DIR, "_sources")
BACKUP_DIR = os.path.join(SOURCE_INFO_DIR, "huopai_backups")
MAX_BACKUPS = 3

# 货盘表各 sheet 的表头所在行；ID列可能叫的名字（都不匹配就追加到表尾）
HUOPAI_SHEETS = {"培育钻": 2, "天然钻": 1}
# 货盘表一行 = 一个 SKU(商家编码)，而平台商品ID是"款"级的(同款各规格都一样)，
# 所以除了「产品id」，还要回填每行自己的「规格id」(平台 SKU ID)。
HUOPAI_ID_HEADERS = {
    "wechat": ("微信小店", "微信商品id", "微信小店商品id", "微信产品id"),
    "xhs": ("小红书产品id", "小红书商品id"),
    "wechat_sku": ("微信规格id", "微信sku id", "微信规格id(sku)"),
    "xhs_sku": ("小红书规格id", "小红书sku id", "小红书规格id(sku)"),
}
HUOPAI_PLATFORM_LABEL = {"wechat": "微信小店", "xhs": "小红书产品id",
                         "wechat_sku": "微信规格id", "xhs_sku": "小红书规格id"}
HUOPAI_FIELD_PLATFORM = {"wechat": "wechat", "xhs": "xhs", "wechat_sku": "wechat", "xhs_sku": "xhs"}
# 写入的ID列统一拉到这个宽度: 原表这些列宽只有 6 个字符, 24 位ID会溢出显示到相邻列,
# 看着像"占用了别的平台ID的位置"
HUOPAI_ID_COLUMN_WIDTH = 24


def _source_info_path(batch_id):
    return os.path.join(SOURCE_INFO_DIR, f"{batch_id}.json")


def save_source_info(batch_id, path, filename):
    """记下批次的来源文件(只有货盘直读才有 path), 供「回填ID到货盘表」定位原文件。"""
    try:
        os.makedirs(SOURCE_INFO_DIR, exist_ok=True)
        with open(_source_info_path(batch_id), "w", encoding="utf-8") as handle:
            json.dump({"path": os.path.abspath(path) if path else "", "filename": filename or ""},
                      handle, ensure_ascii=False)
    except Exception as exc:      # 记不下来不影响导入
        print(f"[source-info] 记录来源失败: {exc}")


def load_source_info(batch_id):
    path = _source_info_path(batch_id)
    if not os.path.isfile(path):
        return {}
    try:
        with open(path, encoding="utf-8") as handle:
            return json.load(handle) or {}
    except Exception:
        return {}


def _col_index(letter):
    index = 0
    for char in str(letter).upper():
        if "A" <= char <= "Z":
            index = index * 26 + (ord(char) - 64)
    return index


def _col_letter(index):
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _xml_text(value):
    return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _patch_row_cells(row_xml, row_number, cells):
    """在一行的 XML 里设置若干单元格: 已有就换内容(保留原样式), 没有就按列序插进去。"""
    for letter in sorted(cells, key=_col_index):
        target = f"{letter}{row_number}"
        cell_re = re.compile(r'<c\b[^>]*\br="%s"[^>]*(?:/>|>.*?</c>)' % re.escape(target), re.S)
        found = cell_re.search(row_xml)
        style = ""
        if found:
            style_match = re.search(r'\bs="[^"]*"', found.group(0))
            style = f" {style_match.group(0)}" if style_match else ""
        new_cell = f'<c r="{target}"{style} t="inlineStr"><is><t>{_xml_text(cells[letter])}</t></is></c>'
        if found:
            row_xml = row_xml[:found.start()] + new_cell + row_xml[found.end():]
            continue
        insert_at = row_xml.rindex("</row>") if "</row>" in row_xml else len(row_xml)
        for cell in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', row_xml):
            if _col_index(cell.group(1)) > _col_index(letter):
                insert_at = cell.start()
                break
        row_xml = row_xml[:insert_at] + new_cell + row_xml[insert_at:]
    # 追加列后把行的 spans(如 spans="1:35") 一起放宽: 有些阅读器(含 openpyxl 只读模式)
    # 会按 spans/dimension 截断, 不改就会出现"写进去了但读出来是空"
    widest = max(_col_index(letter) for letter in cells)
    spans = re.search(r'spans="(\d+):(\d+)"', row_xml)
    if spans and widest > int(spans.group(2)):
        row_xml = row_xml[:spans.start()] + f'spans="{spans.group(1)}:{widest}"' + row_xml[spans.end():]
    return row_xml


def _unhide_columns(sheet_xml, letters, width=HUOPAI_ID_COLUMN_WIDTH):
    """把指定列由隐藏改成显示, 并把列宽拉到放得下 24 位ID。

    货盘表里 D~H 是整段隐藏的(小红书产品id/京东/抖音/微信小店都是隐藏列), 而且宽度只有 6 个字符:
    运营即使取消隐藏, ID 也会因居中溢出显示到相邻列上, 看着像"占用了别的平台ID的位置"。
    所以写入的列统一放开 + 加宽, 其余列(含同一段里其它隐藏列)原样不动。
    <col min="4" max="8" hidden="1"/> 这种区间会被拆开, 只有目标列改属性。
    """
    for letter in sorted(set(letters), key=_col_index):
        index = _col_index(letter)
        target = None
        for match in list(re.finditer(r"<col\b[^>]*/>", sheet_xml)):
            low = re.search(r'min="(\d+)"', match.group(0))
            high = re.search(r'max="(\d+)"', match.group(0))
            if low and high and int(low.group(1)) <= index <= int(high.group(1)):
                target = match
                break
        if target is None:
            # 该列没有列定义(追加在表尾的新列就是这样): 补一条, 保持 min 升序
            definition = f'<col min="{index}" max="{index}" width="{width}" customWidth="1"/>'
            if "</cols>" in sheet_xml:
                following = [(int(m.group(1)), m.start()) for m in re.finditer(r'<col\b[^>]*min="(\d+)"', sheet_xml)]
                later = [position for value, position in following if value > index]
                position = min(later) if later else sheet_xml.rindex("</cols>")
                sheet_xml = sheet_xml[:position] + definition + sheet_xml[position:]
            elif "<sheetData" in sheet_xml:
                sheet_xml = sheet_xml.replace("<sheetData", f"<cols>{definition}</cols><sheetData", 1)
            continue
        attributes = target.group(0)
        low = re.search(r'min="(\d+)"', attributes)
        high = re.search(r'max="(\d+)"', attributes)
        first, last = int(low.group(1)), int(high.group(1))
        body = attributes[len("<col"):-2]
        shown = re.sub(r'\s*(min="\d+"|max="\d+"|hidden="1"|width="[^"]*"|customWidth="[^"]*")', "", body)
        shown += f' width="{width}" customWidth="1"'
        still_hidden = re.sub(r'\s*(min="\d+"|max="\d+")', "", body)      # 其余列: 属性照旧
        pieces = []
        if first < index:
            pieces.append(f'<col min="{first}" max="{index - 1}"{still_hidden}/>')
        pieces.append(f'<col min="{index}" max="{index}"{shown}/>')
        if last > index:
            pieces.append(f'<col min="{index + 1}" max="{last}"{still_hidden}/>')
        sheet_xml = sheet_xml[:target.start()] + "".join(pieces) + sheet_xml[target.end():]
    return sheet_xml


def _patch_sheet_xml(sheet_xml, updates):
    """updates = {行号: {列字母: 值}}；只替换这些行, 其它内容(图片、样式)原样保留。"""
    for row_number in sorted(updates):
        row_re = re.compile(r'<row\b[^>]*\br="%d"[^>]*>.*?</row>' % row_number, re.S)
        match = row_re.search(sheet_xml)
        if not match:
            continue
        patched = _patch_row_cells(match.group(0), row_number, updates[row_number])
        sheet_xml = sheet_xml[:match.start()] + patched + sheet_xml[match.end():]
    # 写入的列如果原本是隐藏的, 一并显示出来, 否则运营看不到回填结果
    sheet_xml = _unhide_columns(sheet_xml, (letter for cells in updates.values() for letter in cells))
    # 新加的列要同步进 <dimension>, 否则阅读器按旧范围截断
    widest = max(_col_index(letter) for cells in updates.values() for letter in cells)
    dimension = re.search(r'<dimension ref="([^"]*)"\s*/>', sheet_xml)
    if dimension:
        parts = dimension.group(1).split(":")
        end = parts[-1]
        end_column = re.match(r"([A-Z]+)(\d*)", end)
        if end_column and _col_index(end_column.group(1)) < widest:
            end = f"{_col_letter(widest)}{end_column.group(2)}"
            ref = f"{parts[0]}:{end}" if len(parts) > 1 else end
            sheet_xml = sheet_xml[:dimension.start()] + f'<dimension ref="{ref}"/>' + sheet_xml[dimension.end():]
    return sheet_xml


def _sheet_xml_targets(archive):
    """sheet 名 → zip 内的 XML 路径（经 workbook.xml + rels 解析, 不靠顺序猜）。"""
    workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
    rels_xml = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rels = dict(re.findall(r'Id="([^"]+)"[^>]*?Target="([^"]+)"', rels_xml))
    targets = {}
    for attributes in re.findall(r'<sheet\b([^>]*?)/?>', workbook_xml):
        name = re.search(r'name="([^"]*)"', attributes)
        rid = re.search(r'r:id="([^"]*)"', attributes)
        if not name or not rid:
            continue
        target = rels.get(rid.group(1), "")
        if not target:
            continue
        target = target[1:] if target.startswith("/") else target
        targets[name.group(1)] = target if target.startswith("xl/") else f"xl/{target}"
    return targets


def _write_workbook(source_path, target_path, updates_by_sheet):
    """把各 sheet 的单元格更新写进 xlsx: 逐条拷贝 zip 内容, 只替换目标 sheet 的 XML。

    ⚠️ 不能用 openpyxl 重新保存整本工作簿 —— 货盘表里的 WPS 内嵌图片(DISPIMG)会被丢掉。
    """
    with zipfile.ZipFile(source_path) as archive:
        targets = _sheet_xml_targets(archive)
        xml_paths = {targets[name]: updates for name, updates in updates_by_sheet.items() if name in targets}
        with zipfile.ZipFile(target_path, "w", zipfile.ZIP_DEFLATED) as output:
            for item in archive.infolist():
                data = archive.read(item.filename)
                if item.filename in xml_paths:
                    data = _patch_sheet_xml(data.decode("utf-8"), xml_paths[item.filename]).encode("utf-8")
                output.writestr(item, data)


def _backup_source(path):
    """备份原货盘表, 只保留最近几份(货盘表 60MB+, 不能无限留)。"""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    stem, extension = os.path.splitext(os.path.basename(path))
    backup = os.path.join(BACKUP_DIR, f"{stem}-{stamp}{extension}")
    shutil.copy2(path, backup)
    backups = sorted(
        (os.path.join(BACKUP_DIR, name) for name in os.listdir(BACKUP_DIR) if name.endswith(extension)),
        key=os.path.getmtime, reverse=True,
    )
    for stale in backups[MAX_BACKUPS:]:
        try:
            os.remove(stale)
        except OSError:
            pass
    return backup


def _verify_written(path, sheet_name, updates):
    """重新读回写入的单元格, 确认真的写进去了(避免写坏 60MB 的货盘表还替换原文件)。"""
    import openpyxl
    lines = sorted(line for line in updates if line > 1)
    if not lines:
        return ""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        for index, values in enumerate(sheet.iter_rows(min_row=1, max_row=lines[-1], values_only=True), start=1):
            if index not in updates:
                continue
            for letter, value in updates[index].items():
                position = _col_index(letter) - 1
                actual = str(values[position] or "").strip() if position < len(values) else ""
                if actual != str(value).strip():
                    return f"{sheet_name}!{letter}{index} 期望 {value} 实际 {actual}"
    finally:
        workbook.close()
    return ""


def _huopai_id_columns(sheet_name, sheet):
    """(表头行, 表头内容, {字段: 列字母}, 需要追加的表头{字段: 名称})。

    字段有 4 个: wechat/xhs = 款级「产品id」列, wechat_sku/xhs_sku = 行级「规格id」列。
    """
    header_row = HUOPAI_SHEETS[sheet_name]
    headers = []
    for row in sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = [str(cell or "").strip() for cell in row]
    normalized = [re.sub(r"[\s()（）]+", "", name).lower() for name in headers]
    next_index = max(sheet.max_column or 0, len(headers))
    columns, appended = {}, {}
    for field, names in HUOPAI_ID_HEADERS.items():
        wanted = {re.sub(r"[\s()（）]+", "", name).lower() for name in names}
        index = next((i for i, name in enumerate(normalized) if name in wanted), None)
        if index is None:
            next_index += 1
            index = next_index - 1
            appended[field] = HUOPAI_PLATFORM_LABEL[field]
        columns[field] = _col_letter(index + 1)
    return header_row, headers, columns, appended


def _sku_key(value):
    return re.sub(r"\s+", "", str(value or "")).upper()


def build_huopai_updates(source_path, rows, ids, sku_ids=None):
    """算出要写进货盘表的单元格: ({sheet: {行: {列: 值}}}, 统计)。

    - 产品id(款级): 同一款所有行一样, 来自发布记录(ids)；
    - 规格id(行级): 每个 SKU 一行不一样, 来自平台商品详情(sku_ids = {平台: {商品编码: {商家编码: 规格id}}})。
    只写"当前为空"的单元格；已有值的行只统计不动(人工填过的不覆盖)；
    导入后货盘表被改过(商家编码对不上)的行直接跳过, 宁可不写也不写错行。
    """
    import openpyxl
    sku_ids = sku_ids or {}
    updates, stats = {}, {"filled": 0, "existing": 0, "mismatch": 0, "sku_missing": 0,
                          "filled_product": 0, "filled_sku": 0, "sheets": {}}
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    try:
        for sheet_name in HUOPAI_SHEETS:
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            header_row, headers, columns, appended = _huopai_id_columns(sheet_name, sheet)
            code_index = next((i for i, name in enumerate(headers) if "商家编码" in name), None)
            wanted = {}
            for row in rows:
                if str(row.get("_源表") or "") != sheet_name:
                    continue
                try:
                    line = int(row.get("_源行") or 0)
                except (TypeError, ValueError):
                    continue
                if line <= header_row:
                    continue
                cells = wanted.setdefault(line, {})
                for field, letter in columns.items():
                    platform = HUOPAI_FIELD_PLATFORM[field]
                    if field.endswith("_sku"):
                        # 规格id按"这一行的商家编码"取: 优先用发品时真正提交的 SKU编码,
                        # 再退回原表里的原始编码(老数据/清洗过的编码靠它兜底)
                        variants = sku_ids.get(platform, {}).get(row.get("product_code"), {})
                        value = None
                        for key in (row.get("sku_code"), row.get("_源码")):
                            value = variants.get(_sku_key(key))
                            if value:
                                break
                        if not value and (ids.get(platform) or {}).get(row.get("product_code")):
                            stats["sku_missing"] += 1
                    else:
                        value = (ids.get(platform) or {}).get(row.get("product_code"))
                    if value:
                        cells[letter] = str(value)
                if cells:
                    cells["__code"] = str(row.get("_源码") or "")
            if appended:
                header_cells = wanted.setdefault(header_row, {})
                for field, label in appended.items():
                    header_cells[columns[field]] = label
            if not wanted:
                continue
            last_line = max(wanted)
            sheet_updates = {}
            for index, values in enumerate(sheet.iter_rows(min_row=1, max_row=last_line, values_only=True), start=1):
                cells = wanted.get(index)
                if not cells:
                    continue
                if index == header_row:
                    sheet_updates[index] = dict(cells)
                    continue
                record = str(cells.pop("__code", "") or "")
                actual_code = str(values[code_index] or "").strip() if code_index is not None and code_index < len(values) else ""
                if record and actual_code != record:
                    stats["mismatch"] += 1
                    continue
                for letter, value in list(cells.items()):
                    position = _col_index(letter) - 1
                    current = str(values[position] or "").strip() if position < len(values) else ""
                    if current:
                        stats["existing"] += 1
                        cells.pop(letter)
                        continue
                    stats["filled"] += 1
                    if _field_of_column(columns, letter).endswith("_sku"):
                        stats["filled_sku"] += 1
                    else:
                        stats["filled_product"] += 1
                if cells:
                    sheet_updates[index] = cells
            if sheet_updates:
                updates[sheet_name] = sheet_updates
                stats["sheets"][sheet_name] = {
                    "rows": len([line for line in sheet_updates if line != header_row]),
                    "columns": columns, "appended": appended,
                }
    finally:
        workbook.close()
    return updates, stats


def _field_of_column(columns, letter):
    return next((field for field, value in columns.items() if value == letter), "")


def latest_platform_ids(shop_id=None):
    """商品编码 → {平台: 平台商品ID}，取最近一次发布成功的记录。

    为什么不只看"当前任务": 运营可能分几次发(分批发布/换批次重新导入/重试),
    货盘表里要登记的是"这个商品在平台上是几号", 所以按商品取最近一次成功记录 ——
    与「发布状态」页「已发布/未发布」的口径一致(它也是按商品而不是按批次汇总)。
    同一商品同一平台存在多个不同ID(重发过、或跨店发布到两家店)时: 优先当前店, 仍不唯一则
    记进 ambiguous, 由前端提示运营确认(仍写入最近一次, 避免整列空着)。
    """
    conn = db()
    rows = conn.execute(
        "SELECT product_code, platform, platform_product_id, shop_id, finished_at FROM publish_items "
        "WHERE platform_product_id IS NOT NULL AND platform_product_id<>'' "
        "AND status IN ('success','partial') ORDER BY finished_at"
    ).fetchall()
    conn.close()
    candidates = {}
    for row in rows:
        key = (row["platform"], row["product_code"])
        bucket = candidates.setdefault(key, [])
        if not any(item["id"] == str(row["platform_product_id"]) for item in bucket):
            bucket.append({"id": str(row["platform_product_id"]), "shop": row["shop_id"], "at": row["finished_at"] or ""})
    ids, shops, ambiguous = {}, {}, {}
    for (platform, code), bucket in candidates.items():
        same_shop = [item for item in bucket if item["shop"] == shop_id]
        pool = same_shop or bucket
        chosen = sorted(pool, key=lambda item: item["at"])[-1]
        ids.setdefault(platform, {})[code] = chosen["id"]
        # 记下这个ID是哪家店发的: 查规格ID必须用同一家店(跨店发布时 A 店查不到 B 店的商品)
        shops.setdefault(platform, {})[code] = chosen["shop"]
        if len({item["id"] for item in pool}) > 1:
            ambiguous.setdefault(platform, []).append(code)
    return ids, shops, ambiguous


def platform_sku_ids(platform, product_id, shop_id=None):
    """查平台商品详情, 拿「商家编码 → 规格(SKU)ID」对应关系。

    货盘表一行 = 一个 SKU(商家编码), 但平台给的商品ID是"款"级的、同款各规格完全一样,
    所以行级要填的规格ID只能回平台查:
      微信   GET /products/{pid}   → product.skus[].{sku_id, sku_code}
      小红书 GET /items/{item_id}  → result.skuInfos[].{id, erpCode}
    查不到就抛异常(由调用方记进"查不到"清单, 不影响产品ID的回填)。
    """
    target = urllib.parse.quote(str(product_id))
    if platform == "wechat":
        url = f"{WECHAT_API_BASE}/products/{target}"
    elif platform == "xhs":
        url = f"{XHS_API_BASE}/items/{target}"
    else:
        raise RuntimeError(f"不支持的平台：{platform}")
    data = _http_get_json(url, timeout=60, shop_id=shop_id)
    mapping = {}
    if platform == "wechat":
        product = (data.get("result") or {}).get("product") or {}
        for sku in product.get("skus") or []:
            sku_id = str(sku.get("sku_id") or "").strip()
            for key in ("out_sku_id", "sku_code"):
                code = _sku_key(sku.get(key))
                if code and sku_id:
                    mapping.setdefault(code, sku_id)
    else:
        if data.get("ok") is False:
            raise RuntimeError(str(data.get("detail") or data)[:200])
        for sku in (data.get("result") or {}).get("skuInfos") or []:
            code = _sku_key(sku.get("erpCode"))
            sku_id = str(sku.get("id") or sku.get("skuId") or "").strip()
            if code and sku_id:
                mapping[code] = sku_id
    return mapping


class WritebackBody(BaseModel):
    dry_run: bool = False


@app.post("/jobs/{job_id}/writeback-huopai")
def writeback_huopai(job_id: str, body: WritebackBody | None = None, request: Request = None):
    """【回填商品ID】把平台商品ID/规格ID写进货盘表原文件。

    - 「微信小店」「小红书产品id」= 款级商品ID(同款各规格一样)；
    - 「微信规格id」「小红书规格id」= 行级规格ID, 现查平台商品详情拿(同款不同规格ID不同)；
    dry_run=true 只统计将要写哪些行(不落盘), 前端据此弹确认框；
    正式写入前备份原文件, 写完自检通过才替换, 失败则原文件保持不动。
    """
    shop_id = _current_shop_id(request) if request else shop_registry.default_shop()["shop_id"]
    _verify_shop(shop_id)
    dry_run = bool(body.dry_run) if body else False
    conn = db()
    job_row = conn.execute("SELECT * FROM publish_jobs WHERE id=? AND shop_id=?", (job_id, shop_id)).fetchone()
    if not job_row:
        conn.close(); raise HTTPException(404, "发布任务不存在")
    batch = conn.execute("SELECT id, filename, rows_json FROM batches WHERE id=? AND shop_id=?",
                         (job_row["batch_id"], shop_id)).fetchone()
    conn.close()
    if not batch:
        raise HTTPException(404, "批次不存在")
    source = load_source_info(batch["id"])
    source_path = source.get("path") or ""
    if not source_path or not os.path.isfile(source_path):
        raise HTTPException(400, "这批不是从货盘表直读导入的（或原货盘表已被移动），无法回填；请用「导出结果」")
    ids, id_shops, ambiguous = latest_platform_ids(shop_id)
    if not ids:
        raise HTTPException(400, "还没有成功发布过的商品ID，发布成功后再回填")
    try:
        rows = json.loads(batch["rows_json"] or "[]")
    except Exception:
        rows = []
    if not any(row.get("_源表") for row in rows):
        raise HTTPException(400, "这批数据是旧版本导入的（没记录源表行号），请重新导入货盘表后再发布")
    # 规格ID(行级)要现查平台: 只查"这批数据里真的用到、且已发布成功"的商品, 控制请求数
    used_codes = {str(row.get("product_code") or "") for row in rows}
    sku_ids, sku_failed, gone = {}, [], {}
    for platform, mapping in ids.items():
        gone_codes = set()
        for code, product_id in list(mapping.items()):
            if code not in used_codes:
                continue
            try:
                # 必须用"发这个ID的那家店"去查, 否则跨店发布时另一家店查不到该商品(400)
                sku_ids.setdefault(platform, {})[code] = platform_sku_ids(
                    platform, product_id, (id_shops.get(platform) or {}).get(code) or shop_id)
            except urllib.error.HTTPError as exc:
                detail = _http_error_detail(exc)
                markers = _WECHAT_GONE_MARKERS if platform == "wechat" else _XHS_GONE_MARKERS
                if any(marker.lower() in detail.lower() for marker in markers):
                    # 平台明确说这个商品不存在了(后台删过): 死ID不写进货盘表, 免得运营照着找
                    gone_codes.add(code)
                    gone.setdefault(platform, []).append(code)
                else:
                    sku_failed.append(f"{code}: {detail[:120]}")
            except Exception as exc:
                sku_failed.append(f"{code}: {str(exc)[:120]}")
        for code in gone_codes:
            mapping.pop(code, None)
            (id_shops.get(platform) or {}).pop(code, None)
    try:
        updates, stats = build_huopai_updates(source_path, rows, ids, sku_ids)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(500, f"解析货盘表失败：{exc}")
    result = {
        "path": source_path, "filename": source.get("filename") or batch["filename"],
        "filled": stats["filled"], "existing": stats["existing"], "mismatch": stats["mismatch"],
        "filled_product": stats["filled_product"], "filled_sku": stats["filled_sku"],
        "sku_missing": stats["sku_missing"],
        "sheets": stats["sheets"],
        "platform_ids": {platform: len(mapping) for platform, mapping in ids.items()},
        "ambiguous": {platform: list(codes)[:20] for platform, codes in ambiguous.items()},
        "sku_failed": sku_failed[:20],
        "gone": {platform: list(codes)[:20] for platform, codes in gone.items()},
        "dry_run": dry_run,
    }
    if dry_run:
        return {"ok": True, "result": result}
    if not updates:
        result["message"] = "没有需要写入的单元格（可能都已有值，或商品编码在货盘表里对不上）"
        return {"ok": True, "result": result}
    try:
        backup = _backup_source(source_path)
    except Exception as exc:
        raise HTTPException(500, f"备份货盘表失败（原文件未改动）：{exc}")
    # ⚠️ 临时文件必须留 .xlsx 后缀: 下面自检要用 openpyxl 读它, openpyxl 不认 .tmp 会直接抛异常。
    # 放在原文件同目录(同盘符才能 os.replace 原子替换), 失败路径统一删除。
    stem, extension = os.path.splitext(source_path)
    temp_path = f"{stem}.idbackup-{os.getpid()}{extension or '.xlsx'}"
    try:
        try:
            _write_workbook(source_path, temp_path, updates)
        except Exception as exc:
            raise HTTPException(500, f"写入货盘表失败（原文件未改动）：{exc}")
        problem = ""
        for name, cells in updates.items():
            problem = _verify_written(temp_path, name, cells)
            if problem:
                break
        if problem:
            raise HTTPException(500, f"写入后自检不通过（原文件未改动）：{problem}；备份在 {backup}")
        try:
            os.replace(temp_path, source_path)
        except PermissionError:
            raise HTTPException(409, "货盘表正被占用（可能有人用 Excel/WPS 打开着），请关闭后重试")
        except OSError as exc:
            raise HTTPException(500, f"替换货盘表失败（原文件未改动，新内容在 {temp_path}）：{exc}")
    finally:
        if os.path.exists(temp_path):     # 成功时已被 os.replace 移走, 失败时清掉, 不留垃圾文件
            try:
                os.remove(temp_path)
            except OSError:
                pass
    result["backup"] = backup
    result["message"] = f"已回填 {stats['filled']} 个单元格到 {os.path.basename(source_path)}"
    return {"ok": True, "result": result}


@app.get("/image-roots")
def list_image_roots(request: Request = None):
    """【图片根目录】返回当前店铺允许扫描的图片根目录,供前端下拉选择。

    优先返回店铺配置的 image_root；全局 ALLOWED_IMAGE_ROOTS 兜底。
    """
    shop_id = _current_shop_id(request) if request else shop_registry.default_shop()["shop_id"]
    shop_root = shop_registry.image_root_for(shop_id)
    roots = _allowed_roots_for(shop_id)
    default = shop_root or DEFAULT_IMAGE_ROOT
    return {"ok": True, "default": os.path.abspath(default),
            "roots": [{"path": os.path.abspath(root), "available": os.path.isdir(root)}
                      for root in roots]}


@app.post("/batches/{batch_id}/images/scan")
def scan_batch_images(batch_id: str, body: ImageScanBody | None = None, request: Request = None):
    """【直读图片】扫描图片根目录, 按编码把商品图/SKU 图挂到批次上。
    多店铺: 图片根目录从当前店铺的 image_root 取。

    解决两个实际问题:
      ① 多 SKU 商品的图分散在子文件夹(文件夹名就是 SKU 编码), 浏览器一次只能选一个文件夹,
         结果只读到本色那几张图, 其他规格的图全漏;
      ② 一次要发几百个商品, 走浏览器 Base64 上传要几 GB, 根本不现实。
    服务端直接读网络共享, 233 个商品 4589 张图约 4 秒扫完。

    匹配规则见 image_scanner.py:
      - 商品主图   = 商品文件夹根下的本色图(按 主图1 -> 主图2 -> 白底 -> PNG 排序)
      - SKU 图     = 该 SKU 的「主图2」(运营指定), 落到 row["sku_image"]
      - 通用详情图 = 通用通栏目录, 追加到每个商品的详情图
    扫描阶段只记录共享盘路径(懒拷贝), 不占磁盘; 真正发布到某个商品时,
    才读它的图片字节落盘缓存(内容 sha256 去重), 再上传平台拿素材 URL。

    传 dry_run=true 可只匹配不写库, 先看命中率。
    """
    shop_id = _current_shop_id(request) if request else shop_registry.default_shop()["shop_id"]
    _verify_shop(shop_id)
    try:
        import image_scanner
    except Exception as exc:
        raise HTTPException(500, f"图片扫描模块不可用：{exc}")

    batch = batch_or_404(batch_id, shop_id)
    rows = batch["items"]
    shop_image_root = shop_registry.image_root_for(shop_id)
    root = (body.root if body and body.root else None) or shop_image_root or DEFAULT_IMAGE_ROOT
    months = body.months if body else None
    dry_run = bool(body.dry_run) if body else False
    # 路径收敛: root 只能取白名单内的图片根目录(见 GET /image-roots),
    # 否则该接口可被用来枚举服务器任意文件夹的商品图。
    allowed_roots = _allowed_roots_for(shop_id)
    if not _within_allowed_roots(root, allowed_roots) or not os.path.isdir(root):
        raise HTTPException(400, f"图片根目录不可访问或不在允许的图片目录内：{root}"
                                 f"(允许：{'、'.join(allowed_roots)})")

    try:
        scan = image_scanner.scan_image_root(root, months=months)
    except Exception as exc:
        raise HTTPException(500, f"扫描图片失败：{exc}")

    product_main = {}      # 商品编码 -> [共享盘图片引用]    (多 SKU 行共用, 只解析一次)
    sku_image = {}         # SKU 编码  -> 共享盘图片引用
    errors = []
    unmatched_skus = []
    unmatched_products = []   # 连商品级主图都没挂上的商品编码: 发布时会被平台以"至少需要一张主图"拒

    def localize(path):
        # 懒拷贝: 只记共享盘路径, 真正上传平台时才读字节落盘。
        # 扫描 233 个商品秒级完成、不占磁盘; 只发一部分商品就只读一部分图。
        return f"{SHARE_PREFIX}{path}"

    common_refs = []
    for path in scan["common_images"]:
        try:
            common_refs.append(localize(path))
        except Exception as exc:
            errors.append(f"通用详情图 {os.path.basename(path)}: {exc}")

    products_matched = 0
    skus_matched = 0
    for row in rows:
        code = str(row.get("product_code") or "")
        if code not in product_main:
            base = code.split("-", 1)[0]
            # ① 精确 → ② 去系列标签(如 ZSTZ106XL-星悦 → ZSTZ106XL)
            paths = scan["root_images"].get(code) or scan["root_images"].get(base)
            if not paths:
                # ③ 图库文件夹名可能带中文后缀(如 HZL253福运钻 vs 商品编码 HZL253-如愿),
                #    取最长前缀兜底; 仍匹配不到 = 图库里确实没有这个款
                best_key = ""
                for key in scan["root_images"]:
                    if key and (key.startswith(base) or base.startswith(key)) and len(key) > len(best_key):
                        best_key = key
                paths = scan["root_images"].get(best_key) if best_key else None
            references = []
            for path in (paths or [])[:5]:      # 平台主图最多 5 张
                try:
                    references.append(localize(path))
                except Exception as exc:
                    errors.append(f"{os.path.basename(path)}: {exc}")
            product_main[code] = references
            if references:
                products_matched += 1
            elif code:
                unmatched_products.append(code)
        if product_main[code]:
            row["main_images"] = product_main[code]
        if common_refs:
            row["detail_images"] = list(dict.fromkeys((row.get("detail_images") or []) + common_refs))

        sku_code = str(row.get("sku_code") or "")
        if sku_code not in sku_image:
            hit = image_scanner.match_sku_code(sku_code, scan["sku_images"])
            main2 = ((scan["sku_images"].get(hit) or {}).get("main2") if hit else None)
            if main2:
                try:
                    sku_image[sku_code] = localize(main2)
                except Exception as exc:
                    sku_image[sku_code] = ""
                    errors.append(f"SKU图 {os.path.basename(main2)}: {exc}")
            else:
                sku_image[sku_code] = ""
        if sku_image[sku_code]:
            row["sku_image"] = sku_image[sku_code]
            skus_matched += 1
        else:
            unmatched_skus.append(sku_code or f"line_{row.get('line')}")

    if not dry_run:
        conn = db()
        conn.execute("UPDATE batches SET rows_json=?, revision=revision+1 WHERE id=? AND shop_id=?",
                     (json.dumps(rows, ensure_ascii=False), batch_id, shop_id))
        conn.commit()
        conn.close()

    return {
        "ok": True, "batch_id": batch_id, "dry_run": dry_run,
        "root": root, "months": months,
        "folders": scan["folders"], "scanned_images": scan["scanned"],
        "common_detail_images": len(scan["common_images"]),
        "products_matched": products_matched,
        "products_total": len({r.get("product_code") for r in rows}),
        "skus_matched": skus_matched, "skus_total": len(rows),
        "unmatched_sku_count": len(unmatched_skus),
        "unmatched_skus": unmatched_skus[:200],
        "unmatched_product_count": len(unmatched_products),
        "unmatched_products": unmatched_products[:200],
        "lazy_load": True, "errors": errors[:50],
    }


@app.post("/batches/{batch_id}/sku-image")
def upload_sku_image(batch_id: str, body: SkuImageBody, request: Request):
    """给批次里某个 SKU 补规格图(浏览器上传单张图)。

    落盘到 IMAGE_DIR/{批次id}/{sha256}.{ext}, 返回 local:// 引用;
    前端拿到引用后写回该行 sku_image, 再 PUT /import/{id}/items 持久化。
    引用格式必须与 _local_image 的校验一致, 否则预览/上传素材会失败。
    """
    shop_id = _current_shop_id(request)
    batch_or_404(batch_id, shop_id)
    try:
        content = base64.b64decode(body.content_base64 or "")
    except Exception as exc:
        raise HTTPException(400, f"图片内容解析失败：{exc}")
    if not content:
        raise HTTPException(400, "图片内容为空")
    extension = os.path.splitext(body.filename or "")[1].lower().lstrip(".")
    if extension not in {"jpg", "jpeg", "png", "webp", "gif", "bmp"}:
        extension = "jpg"
    digest = hashlib.sha256(content).hexdigest()
    batch_dir = os.path.join(IMAGE_DIR, batch_id)
    os.makedirs(batch_dir, exist_ok=True)
    path = os.path.join(batch_dir, f"{digest}.{extension}")
    if not os.path.exists(path):
        with open(path, "wb") as image_file:
            image_file.write(content)
    return {"ok": True, "result": {"ref": f"local://{batch_id}/{digest}.{extension}"}}


@app.get("/images/preview")
def preview_image(ref: str, shop_id: str = None, request: Request = None):
    """预览批次里的图片(供浏览器显示纯文本路径无法渲染的图)。

    支持两种引用:
      - local://{批次id}/{hash}.{ext}   浏览器上传或懒拷贝落盘后的本地图
      - disk:{共享盘绝对路径}            图片直读扫描写入的引用
    出于安全只允许读取白名单(店铺 image_root + ALLOWED_IMAGE_ROOTS)下的文件, 否则参数可读任意文件。

    多店铺: 这个接口常被 <img src> 直接调用, 而浏览器自己发的请求**带不了**
    X-Shop-Id 自定义头, 所以这里额外接受 ?shop_id= 查询参数(header 仍可用, 且优先用参数)。
    """
    shop_id = ((shop_id or "").strip()
               or (_current_shop_id(request) if request else shop_registry.default_shop()["shop_id"]))
    allowed_roots = _allowed_roots_for(shop_id)
    source = str(ref or "").strip()
    if source.startswith("local://"):
        found = _local_image(source)
        if not found:
            raise HTTPException(404, "图片不存在")
        content = found[1]
    elif is_share_image(source):
        path = os.path.abspath(source[len(SHARE_PREFIX):])
        # 用 commonpath 判断而非 startswith: 否则 '...\images_evil\x.jpg' 会命中 '...\images'
        if not _within_allowed_roots(path, allowed_roots) or not os.path.isfile(path):
            raise HTTPException(404, "图片不存在或不在允许的图片目录内")
        with open(path, "rb") as image_file:
            content = image_file.read()
    else:
        raise HTTPException(400, "不支持的图片引用")
    if not content:
        raise HTTPException(404, "图片内容为空")
    media = "image/png" if source.lower().endswith(".png") else "image/jpeg"
    return StreamingResponse(io.BytesIO(content), media_type=media)


def batch_or_404(batch_id, shop_id=None):
    """按 ID 取批次。若传了 shop_id，则校验归属；未传则不校验（兼容旧调用）。"""
    conn = db()
    if shop_id:
        row = conn.execute("SELECT * FROM batches WHERE id=? AND shop_id=?", (batch_id, shop_id)).fetchone()
    else:
        row = conn.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    conn.close()
    if not row: raise HTTPException(404, "批次不存在")
    data = dict(row); data["items"] = json.loads(data.pop("rows_json")); data["products"] = group_products(data["items"]); data["mappings"] = json.loads(data.pop("mappings_json")); return data


# ------------------------------------------------------------------
# 批次解析缓存:worker 每发一个商品都要拿批次的 rows/mappings,
# 全量 json.loads + group_products 在 800 商品时单次近百毫秒,
# 2N 次(商品×平台)调用就是 O(N²)。这里按 batch_id 缓存解析结果。
#
# 跨进程失效: _invalidate_batch_cache 只清本进程的缓存, 而"改映射"发生在 API
# 进程、"发品"发生在 worker 进程, 只靠它 worker 会一直用旧 mappings
# (实测: 改完运费模板立刻重试, 平台仍报旧模板 id 不存在)。
# 因此缓存里额外记住 batches.revision, 命中后再比一次(单列查询, 亚毫秒),
# revision 变了就回源重解析。TTL 仅作兜底(防止将来某条写路径漏 bump revision)。
# ------------------------------------------------------------------
_BATCH_CACHE: dict = {}
_BATCH_CACHE_TTL = 60.0


def _invalidate_batch_cache(batch_id: str = None) -> None:
    """写接口改完 batches 后调用(只影响本进程);batch_id 为空则全清。"""
    if batch_id:
        _BATCH_CACHE.pop(batch_id, None)
    else:
        _BATCH_CACHE.clear()


def _batch_revision(batch_id: str):
    """读该批次当前 revision(单列查询, 不做 JSON 解析);批次不存在返回 None。"""
    conn = db()
    try:
        row = conn.execute("SELECT revision FROM batches WHERE id=?", (batch_id,)).fetchone()
    finally:
        conn.close()
    return row["revision"] if row else None


def get_batch_cached(batch_id: str) -> dict:
    """命中、未过期、且 revision 没变时复用解析结果,否则回源 batch_or_404 并写缓存。"""
    hit = _BATCH_CACHE.get(batch_id)
    if hit and (time.time() - hit["ts"]) < _BATCH_CACHE_TTL:
        revision = _batch_revision(batch_id)
        if revision is not None and revision == hit["revision"]:
            return hit["data"]
    data = batch_or_404(batch_id)
    _BATCH_CACHE[batch_id] = {"ts": time.time(), "revision": _batch_revision(batch_id), "data": data}
    return data


@app.get("/import/{batch_id}")
def get_batch(batch_id: str, request: Request):
    shop_id = _current_shop_id(request)
    return {"ok": True, "result": batch_or_404(batch_id, shop_id)}


@app.put("/import/{batch_id}/items")
def update_items(batch_id: str, body: ItemsBody, request: Request):
    shop_id = _current_shop_id(request)
    batch = batch_or_404(batch_id, shop_id); rows = body.items
    errors = sum(bool(row.get("errors")) for row in rows)
    conn = db(); conn.execute("UPDATE batches SET rows_json=?,status=?,valid=?,errors=?, revision=revision+1 WHERE id=? AND shop_id=?", (json.dumps(rows, ensure_ascii=False), "待校验" if errors else "待发布", len(rows)-errors, errors, batch_id, shop_id)); conn.commit(); conn.close()
    _invalidate_batch_cache(batch_id)
    products = group_products(rows)
    return {"ok": True, "result": {"total": len(rows), "product_count": len(products), "valid": len(rows)-errors, "errors": errors, "product_errors": sum(bool(product["errors"]) for product in products)}}


@app.put("/import/{batch_id}/mappings")
def update_mappings(batch_id: str, body: MappingsBody, request: Request):
    shop_id = _current_shop_id(request)
    batch_or_404(batch_id, shop_id); conn = db(); conn.execute("UPDATE batches SET mappings_json=?, revision=revision+1 WHERE id=? AND shop_id=?", (json.dumps(body.mappings, ensure_ascii=False), batch_id, shop_id)); conn.commit(); conn.close(); _invalidate_batch_cache(batch_id); return {"ok": True, "result": body.mappings}


@app.post("/import/{batch_id}/validate")
def validate(batch_id: str, request: Request):
    shop_id = _current_shop_id(request)
    batch = batch_or_404(batch_id, shop_id); rows = batch["items"]; seen = set()
    for row in rows:
        errors = [e for e in row.get("errors", []) if e not in ("商品编码重复", "SKU编码重复")]
        key = row.get("product_code", "")
        sku = row.get("sku_code", "")
        if key and sku and (key, sku) in seen: errors.append("SKU编码重复")
        seen.add((key, sku)); row["errors"] = errors
    # 复用 update_items 落库; 必须把 request 透传过去, 否则它会拿不到店铺上下文(P1 加 shop_id 后新增的参数)
    return update_items(batch_id, ItemsBody(items=rows), request)


@app.post("/import/{batch_id}/publish")
def publish(batch_id: str, body: PublishBody, request: Request):
    # 这里有两个身份，别混淆：
    #   来源店 = 批次所属店（由 X-Shop-Id 决定），提供平台级映射(类目/属性/SKU规格)；
    #   目标店 = 任务写到哪家店（由 body.target_shop_id 决定），决定用谁的凭证 + 店铺参数发布。
    # 不传 target_shop_id 时两者相同，即改造前的同店发布。
    shop_id = _current_shop_id(request)
    _verify_shop(shop_id)
    operator = _current_operator(request)
    batch = batch_or_404(batch_id, shop_id)
    target_shop_id = (body.target_shop_id or "").strip() or shop_id
    _verify_shop(target_shop_id)
    cross_shop = target_shop_id != shop_id
    if batch["errors"]: raise HTTPException(400, "仍有校验错误，不能发布")
    # 校验所选平台在**目标店**下都配置了（凭证与店铺参数都按目标店取）
    platforms = [p for p in body.platforms if p in ("wechat", "xhs")]
    if not platforms: raise HTTPException(400, "至少选择一个发布平台")
    for p in platforms:
        if not shop_registry.has_platform(target_shop_id, p):
            raise HTTPException(400, f"目标店铺未配置 {p} 平台")
        # 跨店复用来源店的平台级映射，所以来源店必须也是同一平台：
        # 微信批次的类目/属性发不到小红书，硬发必被平台拒。
        if cross_shop and not shop_registry.has_platform(shop_id, p):
            raise HTTPException(400, f"来源店铺未配置 {p} 平台，不能跨店发布该平台")
    # 按 product_codes 过滤(分批发布);不传则全量
    products = batch["products"]
    if body.product_codes:
        code_set = set(body.product_codes)
        products = [p for p in products if p["product_code"] in code_set]
        if not products: raise HTTPException(400, "所选商品均不存在")
    # ⚠️ 店铺必填参数(运费模板/物流方案)没选 → 直接拦在这里, 不建任务:
    #    否则任务建出来、worker 逐个商品报"未配置运费模板", 运营只能看着满屏失败项再重发一轮。
    missing_params = _missing_shop_params(batch["mappings"], products, platforms, body.params or {})
    if missing_params:
        where = "目标店铺参数（跨店发布）" if cross_shop else "「平台映射」这一步"
        raise HTTPException(400, "还没选：" + "；".join(missing_params)
                            + f"。请先在{where}选好，再回来点发布（不会创建任务，也不会产生失败记录）")
    # Idempotency guard: do not create another platform item for a combination
    # that already has a publish record in this batch.
    check_conn = db()
    check_conn.execute("BEGIN IMMEDIATE")
    prior = []
    pending = []
    for product in products:
        for platform in platforms:
            # 幂等按**目标店**判断：发给 A 店成功的记录，不该阻止把同一批货发给 B 店
            row = check_conn.execute("SELECT * FROM publish_items WHERE shop_id=? AND batch_id=? AND product_code=? AND platform=? ORDER BY id DESC LIMIT 1", (target_shop_id, batch_id, product["product_code"], platform)).fetchone()
            # deleted=核对后确认平台商品已被删除, cancelled=人工停掉: 都不算"发过", 必须放行重新发,
            # 否则运营在平台后台删完商品后, 这一批永远卡在"已有发布任务，未重复创建"。
            if row and row["status"] not in ("deleted", "cancelled"):
                prior.append(row)
            else:
                pending.append((product["product_code"], platform))
    if prior and not pending:
        check_conn.commit()
        reusable = next((row for row in reversed(prior) if row["status"] in ("queued", "running")), prior[-1])
        check_conn.close()
        status = reusable["status"]
        message = {"success": "所选商品已发布，未重复创建", "failed": "所选商品已有失败项，请使用重试失败项", "partial": "所选商品存在部分成功项，请勿重复创建"}.get(status, "所选商品已有发布任务，未重复创建")
        return {"ok": True, "job_id": reusable["job_id"], "batch_id": batch_id, "platforms": platforms, "existing": True, "created_count": 0, "skipped_count": len(prior), "target_shop_id": target_shop_id, "cross_shop": cross_shop, "message": message}
    # Keep the write transaction open through job/item insertion.
    conn = check_conn
    job_id = uuid.uuid4().hex
    total = len(pending) if pending else len(products) * len(platforms)
    # 分批发布(带了 product_codes) = 运营明确只要发这几件, 标记优先, 插到队列前面走
    partial = bool(body.product_codes)
    priority = 1 if partial else 0
    # 任务与任务项写到**目标店**名下（worker 据此用目标店凭证发布，图片素材也传目标店）；
    # params_json 存该店私有参数，执行时覆盖批次里的来源店参数。
    # 批次本身仍属来源店：状态更新按来源店，绝不改批次的 shop_id。
    conn.execute("INSERT INTO publish_jobs(id,shop_id,batch_id,platforms_json,status,total,operator,created_at,updated_at,params_json) VALUES(?,?,?,?,?,?,?,?,?,?)", (job_id, target_shop_id, batch_id, json.dumps(platforms), "queued", total, operator, now(), now(), json.dumps(body.params or {})))
    for product_code, platform in (pending or [(product["product_code"], platform) for product in products for platform in platforms]):
        conn.execute("INSERT INTO publish_items(job_id,shop_id,batch_id,product_code,platform,status,priority) VALUES(?,?,?,?,?,?,?)", (job_id, target_shop_id, batch_id, product_code, platform, "queued", priority))
    conn.execute("UPDATE batches SET status=? WHERE id=? AND shop_id=?", ("发布任务已创建", batch_id, shop_id)); conn.commit(); conn.close(); _invalidate_batch_cache(batch_id)
    message = f"已创建发布任务({len(products)} 件商品 × {len(platforms)} 平台)" if partial else "发布任务已入队，后台将持续执行"
    if cross_shop:
        message += f"（目标店铺：{target_shop_id}）"
    return {"ok": True, "job_id": job_id, "batch_id": batch_id, "platforms": platforms, "partial": partial, "product_count": len(products), "target_shop_id": target_shop_id, "cross_shop": cross_shop, "message": message}


@app.get("/publish-status")
def publish_status(request: Request):
    """【发布状态】按商品编码汇总**当前店铺所有批次**的发布记录。

    为什么按商品、而不是按批次:
      运营的常见做法是"重新导入一次货盘表再发", 而每次导入都会生成一个新批次。
      如果按 batch_id 统计, 重新导入后"已经发过"的记录全部消失, 又会显示成 250 件
      未发布, 于是再重复发一遍 —— 平台对重复标题直接拒绝(-5000500 标题不能与现有商品标题重复)。
      商品编码跨批次稳定, 所以这里全局汇总, 换批次也不会丢。

    只统计"有意义的终态": success/partial 表示确实发出去过; queued/running 表示正在发;
    failed 表示发过但失败; cancelled 是被手动停掉的, 不算发过(前端按 none 处理);
    deleted 是核对后发现平台商品已被删除的, 同样按 none 处理(可以重新发)。

    返回: {商品编码: {"platforms": {平台: {状态: 条数}}, "last_success_at": ISO时间}}
    """
    shop_id = _current_shop_id(request)
    conn = db()
    rows = conn.execute("SELECT product_code, platform, status, COUNT(*) n, MAX(finished_at) last_at "
                        "FROM publish_items WHERE shop_id=? GROUP BY product_code, platform, status",
                        (shop_id,)).fetchall()
    conn.close()
    result = {}
    for row in rows:
        entry = result.setdefault(row["product_code"], {"platforms": {}, "last_success_at": None})
        entry["platforms"].setdefault(row["platform"], {})[row["status"]] = row["n"]
        if row["status"] == "success" and row["last_at"]:
            if not entry["last_success_at"] or row["last_at"] > entry["last_success_at"]:
                entry["last_success_at"] = row["last_at"]
    return {"ok": True, "result": result}


# ==================================================================
# 核对发布状态: 本地"已发布"记录 vs 平台真实商品
# ==================================================================
ITEM_EXISTS = "exists"    # 平台还有该商品
ITEM_GONE = "gone"        # 平台明确表示不存在/已删除
ITEM_UNKNOWN = "unknown"  # 查不动(网络/鉴权/其它), 不能据此改状态

# "商品不存在"的判定特征(已对官方文档):
#   微信 获取商品: errcode 10020052 = 商品不存在; product.status=6 = 回收站(后台删除后落到这里)
#   小红书 2026-09-16 实测: 查已删除商品返回
#       {"detail":"[product.getItemInfo] 调用失败: error_code=-5000500 item not exist"}
#     即错误码 -5000500 + 文案 "item not exist"。
#     ⚠️ 这里仍然**只按文案**匹配、不认错误码: -5000500 在微信侧是"标题重复"，
#        疑似是跨平台的通用码；只认数字码可能把"标题重复"误判成"商品已删除"。
#        文案 "item not exist" 已经能精确命中这个场景，更安全。
# 注意别用过于宽泛的词(如裸的"不存在"): 微信的"运费模板不存在""类目不存在"等也会命中,
# 会把在售商品误判成已删除, 所以微信只认自己的错误码和明确提到商品的文案。
# ⚠️ 小红书同理, 而且更要收紧: "not found" 正是 Starlette/FastAPI 默认 404 的原文({"detail":"Not Found"}),
#   一旦把它当删除信号, 权限/参数/路径类报错都会命中 →
#   批量把 success 改 deleted → 前端变"未发布" → 幂等保护放行 → 运营重发 → 平台上出现重复商品。
#   所以这里只认"明确指向商品"的文案; 拿不准一律 unknown(状态不变)。
_WECHAT_GONE_MARKERS = ("10020052", "商品不存在", "商品已删除", "商品在回收站")
_XHS_GONE_MARKERS = (
    "商品不存在", "商品已删除",
    "item not exist", "item_not_exist", "item not found",
)


def _http_error_detail(exc):
    """取 HTTPError 响应体里的 detail —— 两个平台服务都用 FastAPI, 真实错误码在 detail 里。"""
    try:
        raw = exc.read().decode("utf-8")
        return str(json.loads(raw).get("detail") or raw)[:500]
    except Exception:
        return str(exc)[:500]


def platform_item_state(platform, product_id, shop_id=None):
    """核对该商品在平台上还在不在, 返回 (state, 说明)。

    只有平台**明确**说商品不存在/已删除时才返回 gone; 网络超时、连接失败、鉴权过期
    一律返回 unknown 并保留原状态 —— 否则一次网络抖动就会把在售商品标成"未发布",
    诱导运营重复发布(平台会以「标题重复」拒绝), 反而更麻烦。
    """
    target = urllib.parse.quote(str(product_id))
    if platform == "wechat":
        url = f"{WECHAT_API_BASE}/products/{target}"
    elif platform == "xhs":
        url = f"{XHS_API_BASE}/items/{target}"
    else:
        return ITEM_UNKNOWN, f"不支持的平台：{platform}"
    try:
        data = _http_get_json(url, timeout=30, shop_id=shop_id)
    except urllib.error.HTTPError as exc:
        detail = _http_error_detail(exc)
        markers = _WECHAT_GONE_MARKERS if platform == "wechat" else _XHS_GONE_MARKERS
        lowered = str(detail).lower()
        if any(marker.lower() in lowered for marker in markers):
            return ITEM_GONE, detail
        return ITEM_UNKNOWN, detail
    except Exception as exc:
        return ITEM_UNKNOWN, f"查询失败：{exc}"
    if platform == "wechat":
        product = (data.get("result") or {}).get("product") or {}
        if not product.get("product_id"):
            return ITEM_GONE, "平台未返回商品数据"
        if str(product.get("status")) == "6":
            return ITEM_GONE, "商品在平台回收站（后台已删除）"
        return ITEM_EXISTS, ""
    # 小红书: 200 也要看返回体。商品被后台删掉后, 网关很可能回 200 + 空 result 而不是报错,
    # 只看"接口有没有报错"会把已删商品一直当成"还在", 核对等于没做(小红书又是失败/删除的重灾区)。
    # ⚠️ 该判定尚未实测: 请拿一个确实已在后台删掉的小红书 itemId 调一次 GET /items/{id},
    #    把真实返回(是报错? 空 result? 还是别的结构)补记到本注释, 再确认这条判定可靠。
    result = data.get("result")
    if data.get("ok") is False:
        # 200 但业务失败 → 不下结论(保留原状态)
        return ITEM_UNKNOWN, f"平台返回业务失败：{str(data.get('detail') or data.get('message') or data)[:200]}"
    if not result:
        # 空 result: 平台确实没给商品数据 → 判定已删除(把原始返回带进 error, 便于事后复盘)
        return ITEM_GONE, f"平台未返回商品数据：{str(data)[:200]}"
    if not isinstance(result, dict):
        return ITEM_UNKNOWN, f"返回结构无法识别：{str(result)[:200]}"
    if not (result.get("itemInfo") or result.get("itemId") or result.get("item_id") or result.get("id")):
        # 有返回但结构不认识 → 不结论(宁可漏判, 也不要把在售商品误标成未发布)
        return ITEM_UNKNOWN, f"返回结构无法识别：{str(result)[:200]}"
    return ITEM_EXISTS, ""


class VerifyPublishStatusBody(BaseModel):
    product_codes: list[str] | None = None  # 不传=核对全部有成功记录的商品
    mode: str = "verify"                    # verify=按平台核对; reset=人工确认已删除, 直接改回未发布


@app.post("/publish-status/verify")
def verify_publish_status(body: VerifyPublishStatusBody, request: Request):
    """【核对发布状态】把本地"已发布"记录拿去平台核一遍, 已删掉的改回"未发布"。

    为什么需要: 运营会在平台后台把商品删掉, 但本地 publish_items 还留着 success,
    step3 就一直显示"已发布" —— 既看着不对, 又会撞上发布接口的幂等保护(同一批次
    同商品已有记录就不再创建), 导致删完后想重发却发不出去。

    处理方式: 平台确认不存在的记录状态改成 deleted(前端按"未发布"处理, 且幂等保护放行);
    查不动的保持原状(unknown), 并在返回里提示, 由运营用 mode=reset 人工兜底。
    """
    shop_id = _current_shop_id(request)
    conn = db()
    sql = ("SELECT id, product_code, platform, platform_product_id, finished_at FROM publish_items "
           "WHERE shop_id=? AND status IN ('success','partial') AND COALESCE(platform_product_id,'')<>''")
    params = [shop_id]
    if body.product_codes:
        codes = list({str(code) for code in body.product_codes})
        sql += f" AND product_code IN ({','.join('?' for _ in codes)})"
        params.extend(codes)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    # 同一商品+平台可能有多条历史成功记录: 只拿最新那条的 platform_product_id 去平台核对
    latest = {}
    for row in rows:
        key = (row["product_code"], row["platform"])
        if key not in latest or str(row["finished_at"] or "") > str(latest[key]["finished_at"] or ""):
            latest[key] = row
    targets = list(latest.values())
    if not targets:
        return {"ok": True, "result": {"checked": 0, "deleted": 0, "exists": 0, "unknown": 0, "products": [], "unknown_items": []}}
    states = {}
    if body.mode == "reset":
        # 人工兜底: 运营自己确认这些商品在后台删了, 不再请求平台(也用于小红书查不出的情况)
        for row in targets:
            states[(row["product_code"], row["platform"])] = (ITEM_GONE, "人工标记为已在平台删除")
    else:
        from concurrent.futures import ThreadPoolExecutor, as_completed
        with ThreadPoolExecutor(max_workers=min(6, len(targets))) as pool:
            futures = {pool.submit(platform_item_state, row["platform"], row["platform_product_id"], shop_id): row for row in targets}
            for future in as_completed(futures):
                row = futures[future]
                try:
                    states[(row["product_code"], row["platform"])] = future.result()
                except Exception as exc:
                    states[(row["product_code"], row["platform"])] = (ITEM_UNKNOWN, str(exc)[:200])
    deleted_products, unknown_items, exists_count = set(), [], 0
    conn = db()
    for row in targets:
        state, note = states.get((row["product_code"], row["platform"]), (ITEM_UNKNOWN, ""))
        if state == ITEM_GONE:
            # 只改「核对过的那条身份」: WHERE 必须带上 platform_product_id。
            # 判定用的 platform_product_id 是上面 SELECT 时读到的, 而这里原本按"商品+平台"宽条件更新;
            # 若两步之间运营正好重发成功(新 success 行落库), 这条新行会被一起改成 deleted →
            # 刚发成功的商品显示"未发布" → 再发一次 → 平台上出现重复商品(只能人工删)。
            # 带上 platform_product_id 后, 并发新增的记录平台 ID 必然不同, 不会被误伤;
            # 同一身份的其它历史记录仍一起改, 保证汇总接口按状态计数时不会漏成"已发布"。
            conn.execute("UPDATE publish_items SET status='deleted', error=? "
                         "WHERE shop_id=? AND product_code=? AND platform=? AND platform_product_id=? "
                         "AND status IN ('success','partial')",
                         (f"平台已删除（{now()} 核对，平台ID={row['platform_product_id']}）：{note}"[:500],
                          shop_id, row["product_code"], row["platform"], row["platform_product_id"]))
            deleted_products.add(row["product_code"])
        elif state == ITEM_EXISTS:
            exists_count += 1
        else:
            unknown_items.append({"product_code": row["product_code"], "platform": row["platform"], "reason": note})
    conn.commit()
    conn.close()
    return {"ok": True, "result": {
        "checked": len(targets), "deleted": len(deleted_products), "exists": exists_count,
        "unknown": len(unknown_items), "products": sorted(deleted_products), "unknown_items": unknown_items[:50],
    }}


@app.get("/jobs/{job_id}")
def job(job_id: str, request: Request):
    shop_id = _current_shop_id(request)
    conn = db(); j = conn.execute("SELECT * FROM publish_jobs WHERE id=? AND shop_id=?", (job_id, shop_id)).fetchone()
    if not j: conn.close(); raise HTTPException(404, "发布任务不存在")
    items = [dict(r) for r in conn.execute("SELECT * FROM publish_items WHERE job_id=? AND shop_id=? ORDER BY id", (job_id, shop_id)).fetchall()]
    conn.close()
    return {"ok": True, "result": {"job_id": job_id, "batch_id": j["batch_id"], "status": j["status"], "total": j["total"], "processed": j["processed"], "success": j["success"], "failed": j["failed"], "items": items}}


@app.get("/image-cache/stats")
def image_cache_stats(request: Request):
    """图片上传缓存统计:DB 条目数、内存命中、DB 命中、未命中次数。"""
    shop_id = _current_shop_id(request)
    conn = db()
    total = conn.execute("SELECT COUNT(*) as c FROM image_upload_cache WHERE shop_id=?", (shop_id,)).fetchone()["c"]
    wechat_count = conn.execute("SELECT COUNT(*) as c FROM image_upload_cache WHERE shop_id=? AND platform='wechat'", (shop_id,)).fetchone()["c"]
    xhs_count = conn.execute("SELECT COUNT(*) as c FROM image_upload_cache WHERE shop_id=? AND platform='xhs'", (shop_id,)).fetchone()["c"]
    conn.close()
    total_lookups = _image_cache_hit_mem + _image_cache_hit_db + _image_cache_miss
    hit_rate = round((_image_cache_hit_mem + _image_cache_hit_db) / total_lookups * 100, 1) if total_lookups else 0
    return {"ok": True, "result": {
        "db_total": total, "db_wechat": wechat_count, "db_xhs": xhs_count,
        "mem_size": len(_IMAGE_CACHE_MEM), "mem_max": _IMAGE_CACHE_MAX,
        "hit_mem": _image_cache_hit_mem, "hit_db": _image_cache_hit_db,
        "miss": _image_cache_miss, "total_lookups": total_lookups, "hit_rate_pct": hit_rate,
    }}


@app.delete("/image-cache")
def clear_image_cache(request: Request):
    """清空当前店铺的图片上传缓存(内存 + SQLite)。用于平台素材被删、强制重传等场景。"""
    shop_id = _current_shop_id(request)
    global _image_cache_hit_mem, _image_cache_hit_db, _image_cache_miss
    _IMAGE_CACHE_MEM.clear()
    _IMAGE_CACHE_ORDER.clear()
    _image_cache_hit_mem = _image_cache_hit_db = _image_cache_miss = 0
    conn = db(); conn.execute("DELETE FROM image_upload_cache WHERE shop_id=?", (shop_id,)); conn.commit(); conn.close()
    return {"ok": True, "message": "图片上传缓存已清空"}


# ==================================================================
# 类目别名映射表:内部类目 -> 微信/小红书平台类目。
# 运营人工确认一次类目组后存入此表,后续批次自动命中,不再重复选。
# ==================================================================
@app.post("/jobs/{job_id}/retry")
def retry_job(job_id: str, body: RetryBody, request: Request):
    """仅重新排队失败项；部分成功项不自动重试，避免重复创建平台商品。"""
    shop_id = _current_shop_id(request)
    conn = db()
    if not conn.execute("SELECT id FROM publish_jobs WHERE id=? AND shop_id=?", (job_id, shop_id)).fetchone():
        conn.close()
        raise HTTPException(404, "发布任务不存在")
    requested = set(body.item_ids or [])
    if requested:
        placeholders = ",".join("?" for _ in requested)
        rows = conn.execute(f"SELECT * FROM publish_items WHERE shop_id=? AND job_id=? AND id IN ({placeholders})", [shop_id, job_id, *requested]).fetchall()
        if requested - {row["id"] for row in rows}:
            conn.close()
            raise HTTPException(400, "只能重试当前任务中的商品")
    else:
        rows = conn.execute("SELECT * FROM publish_items WHERE shop_id=? AND job_id=? AND status='failed'", (shop_id, job_id)).fetchall()
    conn.execute("BEGIN IMMEDIATE")
    retried = skipped = 0
    for row in rows:
        if row["status"] == "failed":
            conn.execute("UPDATE publish_items SET status='queued', error=NULL, started_at=NULL, finished_at=NULL WHERE id=? AND shop_id=?", (row["id"], shop_id))
            retried += 1
        elif row["status"] == "partial":
            skipped += 1
    if retried:
        conn.execute("UPDATE publish_jobs SET status='queued', processed=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status IN ('success','partial')), success=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status='success'), failed=(SELECT COUNT(*) FROM publish_items WHERE shop_id=? AND job_id=? AND status='partial'), updated_at=? WHERE id=? AND shop_id=?", (shop_id, job_id, shop_id, job_id, shop_id, job_id, now(), job_id, shop_id))
    conn.commit(); conn.close()
    message = f"已将 {retried} 个失败商品重新排队" if retried else ("部分成功项不会自动重试，以免重复创建平台商品" if skipped else "当前没有可重试的失败商品")
    return {"ok": True, "job_id": job_id, "retried": retried, "skipped": skipped, "message": message}


@app.get("/category-aliases")
def list_aliases(request: Request):
    """类目别名映射表（跨店铺共享读取）。

    ⚠️ 这里刻意返回**所有店铺**的别名，而不是只看当前店：
       平台类目树是全局的（同一平台下所有店铺共用同一棵树），别名表里也只存
       类目链 / 属性默认值 / 规格映射，不含任何店铺私有参数（运费模板、物流方案、
       品牌都不在这张表里）。原来按 shop_id 过滤，会让"4 家店卖同类珠宝"变成
       同一套映射要配 4 遍，而实际上配一次就该全平台通用。

    同内部类目可能分散在不同店铺的行里（A 店配了微信侧、B 店只配了小红书侧），
    所以两侧各自取最近一次有值的，拼成一份完整映射返回。
    """
    conn = db()
    rows = conn.execute("SELECT * FROM category_aliases ORDER BY updated_at DESC").fetchall()
    conn.close()
    merged: dict = {}
    for r in rows:
        key = r["internal_category"]
        entry = merged.get(key)
        if entry is None:
            entry = merged[key] = {
                "id": r["id"], "internal_category": key, "shop_id": r["shop_id"],
                "wechat": None, "xhs": None, "updated_at": r["updated_at"],
                "shop_ids": [],
            }
        entry["shop_ids"].append(r["shop_id"])
        if entry["wechat"] is None and r["wechat_json"]:
            entry["wechat"] = json.loads(r["wechat_json"])
        if entry["xhs"] is None and r["xhs_json"]:
            entry["xhs"] = json.loads(r["xhs_json"])
    return {"ok": True, "result": list(merged.values())}


@app.post("/category-aliases")
def save_alias(body: AliasBody, request: Request):
    shop_id = _current_shop_id(request)
    key = (body.internal_category or "").strip()
    if not key:
        raise HTTPException(400, "internal_category 不能为空")
    conn = db()
    existing = conn.execute("SELECT wechat_json, xhs_json FROM category_aliases WHERE shop_id=? AND internal_category=?", (shop_id, key)).fetchone()
    # 合并语义:本次未传的一侧保留旧值(微信和小红书经常分两次确认)
    wechat = body.wechat if body.wechat is not None else (json.loads(existing["wechat_json"]) if existing and existing["wechat_json"] else None)
    xhs = body.xhs if body.xhs is not None else (json.loads(existing["xhs_json"]) if existing and existing["xhs_json"] else None)
    if not wechat and not xhs:
        conn.close()
        raise HTTPException(400, "wechat 和 xhs 至少提供一个")
    conn.execute(
        """INSERT INTO category_aliases(shop_id, internal_category, wechat_json, xhs_json, updated_at)
           VALUES(?,?,?,?,?)
           ON CONFLICT(shop_id, internal_category) DO UPDATE SET
             wechat_json=excluded.wechat_json, xhs_json=excluded.xhs_json, updated_at=excluded.updated_at""",
        (shop_id, key, json.dumps(wechat, ensure_ascii=False) if wechat else None,
         json.dumps(xhs, ensure_ascii=False) if xhs else None, now()))
    conn.commit(); conn.close()
    return {"ok": True, "result": {"internal_category": key, "wechat": wechat, "xhs": xhs}}


@app.delete("/category-aliases/{alias_id}")
def delete_alias(alias_id: int, request: Request):
    """删除一条类目映射。

    ⚠️ 列表是跨店合并展示的，删除也必须按「内部类目」把所有店铺的同名行一起删掉，
       否则会出现"删了它、切到另一个店又冒出来"的诡异现象。
    """
    conn = db()
    row = conn.execute("SELECT internal_category FROM category_aliases WHERE id=?", (alias_id,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "类目映射不存在")
    conn.execute("DELETE FROM category_aliases WHERE internal_category=?", (row["internal_category"],))
    conn.commit(); conn.close()
    return {"ok": True}


# ------------------------------------------------------------------
# 映射表批量导入:运营提前整理"内部类目 -> 平台类目路径"Excel,
# 导入时按官方类目名逐级反查平台类目ID并校验,成功行存入映射表。
# ------------------------------------------------------------------
# ⚠️ 默认值必须是 127.0.0.1 而不是 localhost：Windows 下 localhost 优先解析成 IPv6 [::1]，
#    而两个平台服务只监听 IPv4，导致每次调用都要先把 ::1 连超时才回退 IPv4
#    （实测 worker 进程里堆着一排 SYN_SENT 到 [::1]:8000/8010，明显拖慢任务领取与图片上传）。
#    127.0.0.1 是纯 IPv4，行为确定。
WECHAT_API_BASE = os.environ.get("WECHAT_API_BASE", "http://127.0.0.1:8000")
XHS_API_BASE = os.environ.get("XHS_API_BASE", "http://127.0.0.1:8010")


def _http_get_json(url, timeout=60, shop_id=None, extra_headers=None):
    headers = {}
    if shop_id:
        headers[SHOP_HEADER] = shop_id
    if extra_headers:
        headers.update(extra_headers)
    if headers:
        req = urllib.request.Request(url, headers=headers)
    else:
        req = url
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def split_category_path(text):
    return [part.strip() for part in re.split(r"[>＞]", str(text or "")) if part.strip()]


def parse_mapping_sheet(raw, filename):
    """解析映射表文件,返回 [{内部类目/微信类目路径/小红书类目路径...}] 原始行。"""
    extension = os.path.splitext(filename or "")[1].lower()
    if extension == ".xls":
        raise HTTPException(400, "暂不支持旧版 .xls，请另存为 .xlsx 后再导入")
    if extension == ".csv":
        try:
            return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
        except UnicodeDecodeError:
            raise HTTPException(400, "CSV 请保存为 UTF-8 编码后再导入")
        except Exception as exc:
            raise HTTPException(400, f"CSV 解析失败：{exc}")
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(x or "").strip() for x in rows[0]]
        result = []
        for cells in rows[1:]:
            if not any(cell not in (None, "") for cell in cells):
                continue
            result.append({headers[i]: (cells[i] if i < len(cells) else "") for i in range(len(headers)) if headers[i]})
        return result
    except Exception as exc:
        raise HTTPException(400, f"Excel 解析失败：{exc}")


def resolve_wechat_path(parts, cache, shop_id=None):
    """按官方名称路径整链反查微信类目,返回 (结果, 错误原因)。"""
    keyword = parts[-1]
    if keyword not in cache:
        try:
            data = _http_get_json(f"{WECHAT_API_BASE}/categories?keyword={urllib.parse.quote(keyword)}", timeout=180, shop_id=shop_id)
        except Exception as exc:
            raise HTTPException(502, f"微信类目服务({WECHAT_API_BASE})不可用："
                                     f"{_http_error_detail(exc) if isinstance(exc, urllib.error.HTTPError) else exc}")
        cache[keyword] = [r for r in data.get("results", []) if r.get("leaf")]
    leaves = cache[keyword]
    matched = [r for r in leaves if [str(n.get("name", "")).strip() for n in r.get("chain", [])] == parts]
    if len(matched) == 1:
        return {"category": matched[0]["path"], "chain": matched[0].get("chain", [])}, None
    candidates = "；".join(r.get("path", "") for r in leaves[:5]) or "无"
    return None, f"未匹配到唯一微信类目（含“{keyword}”的候选：{candidates}）"


def resolve_xhs_path(parts, cache, shop_id=None):
    """按官方名称路径逐级反查小红书类目,末级必须是叶子,返回 (结果, 错误原因)。"""
    parent_id = None
    chain = []
    last = None
    for index, part in enumerate(parts):
        key = str(parent_id or "")
        if key not in cache:
            url = f"{XHS_API_BASE}/categories" + (f"?category_id={urllib.parse.quote(key)}" if parent_id else "")
            try:
                data = _http_get_json(url, timeout=60, shop_id=shop_id)
            except Exception as exc:
                raise HTTPException(502, f"小红书类目服务({XHS_API_BASE})不可用："
                                         f"{_http_error_detail(exc) if isinstance(exc, urllib.error.HTTPError) else exc}")
            result = data.get("result") or []
            cache[key] = result if isinstance(result, list) else []
        options = cache[key]
        matches = [c for c in options if str(c.get("name", "")).strip() == part]
        if not matches:
            names = "、".join(str(c.get("name", "")) for c in options[:10]) or "空"
            return None, f"第{index + 1}级没有“{part}”，该级可选：{names}"
        last = matches[0]
        parent_id = str(last.get("id") or last.get("categoryId") or "")
        chain.append({"id": parent_id, "name": str(last.get("name", "")).strip()})
    if not (last.get("isLeaf") or last.get("leaf")):
        return None, f"“{parts[-1]}”不是叶子类目，请写到平台最后一级"
    return {"category": " > ".join(c["name"] for c in chain), "chain": chain, "category_id": chain[-1]["id"]}, None


@app.get("/category-aliases/template")
def alias_template():
    headers = ["内部类目", "微信类目路径", "小红书类目路径"]
    try:
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "类目映射"
        sheet.append(headers)
        sheet.append(["钻石 > 钻石首饰 > 项链", "珠宝首饰 > 天然钻石 > 项链", "黄金/彩宝/钻石/珍珠 > 钻石首饰 > 颈饰/挂件/吊坠"])
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        sheet.freeze_panes = "A2"
        note = workbook.create_sheet("填写说明")
        for line in [
            "1. 内部类目必须与商品导入表里的内部类目完全一致，同名行会覆盖更新已有映射。",
            "2. 平台类目路径必须使用平台官方类目名，各级之间用 > 分隔，写到叶子级为止。",
            "3. 只填一侧也可以，未填写的一侧保留映射表里已有的旧值。",
            "4. 导入时系统自动把名称反查成平台类目ID并逐级校验，失败行会给出原因，修正后重新上传即可。",
        ]:
            note.append([line])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(iter([output.read()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=category-alias-template.xlsx"})
    except Exception as exc:
        raise HTTPException(500, f"生成映射表模板失败：{exc}")


@app.post("/category-aliases/import")
def import_aliases(body: AliasImportBody, request: Request):
    shop_id = _current_shop_id(request)
    try:
        raw = base64.b64decode(body.content_base64)
    except Exception:
        raise HTTPException(400, "文件编码无效")
    rows = parse_mapping_sheet(raw, body.filename)
    wechat_cache, xhs_cache, report, resolved = {}, {}, [], []
    imported = failed = skipped = 0
    for index, row in enumerate(rows, 2):
        internal = value(row, "内部类目", "类目")
        wechat_path = value(row, "微信类目路径", "微信类目", "微信")
        xhs_path = value(row, "小红书类目路径", "小红书类目", "小红书")
        if not internal:
            skipped += 1
            report.append({"line": index, "internal_category": "", "wechat": "跳过：缺少内部类目", "xhs": ""})
            continue
        if not wechat_path and not xhs_path:
            skipped += 1
            report.append({"line": index, "internal_category": internal, "wechat": "", "xhs": "跳过：两侧类目都为空"})
            continue
        wechat = xhs = None
        wechat_msg = xhs_msg = "未填写"
        if wechat_path:
            wechat, err = resolve_wechat_path(split_category_path(wechat_path), wechat_cache, shop_id)
            wechat_msg = "成功" if wechat else err
        if xhs_path:
            xhs, err = resolve_xhs_path(split_category_path(xhs_path), xhs_cache, shop_id)
            xhs_msg = "成功" if xhs else err
        if wechat or xhs:
            imported += 1
            resolved.append((internal, wechat, xhs))
        else:
            failed += 1
        report.append({"line": index, "internal_category": internal, "wechat": wechat_msg, "xhs": xhs_msg})
    if resolved:
        conn = db()
        for internal, wechat, xhs in resolved:
            # 合并语义与单条保存一致:本次未解析出的一侧保留旧值
            existing = conn.execute("SELECT wechat_json, xhs_json FROM category_aliases WHERE shop_id=? AND internal_category=?", (shop_id, internal)).fetchone()
            merged_wechat = wechat or (json.loads(existing["wechat_json"]) if existing and existing["wechat_json"] else None)
            merged_xhs = xhs or (json.loads(existing["xhs_json"]) if existing and existing["xhs_json"] else None)
            conn.execute(
                """INSERT INTO category_aliases(shop_id, internal_category, wechat_json, xhs_json, updated_at)
                   VALUES(?,?,?,?,?)
                   ON CONFLICT(shop_id, internal_category) DO UPDATE SET
                     wechat_json=excluded.wechat_json, xhs_json=excluded.xhs_json, updated_at=excluded.updated_at""",
                (shop_id, internal, json.dumps(merged_wechat, ensure_ascii=False) if merged_wechat else None,
                 json.dumps(merged_xhs, ensure_ascii=False) if merged_xhs else None, now()))
        conn.commit(); conn.close()
    return {"ok": True, "result": {"imported": imported, "failed": failed, "skipped": skipped, "report": report}}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host=BULK_API_HOST, port=BULK_API_PORT)
