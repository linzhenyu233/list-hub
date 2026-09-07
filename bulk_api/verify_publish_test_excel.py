# -*- coding: utf-8 -*-
"""Verify the generated workbook without creating products on either platform."""

import base64
import json

import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import bulk_api
from generate_publish_test_excel import OUTPUT


BULK = "http://127.0.0.1:8020"
WECHAT = "http://127.0.0.1:8000"
XHS = "http://127.0.0.1:8010"


def result_list(data, *keys):
    value = data.get("result") or data
    if isinstance(value, list):
        return value
    for key in keys:
        if isinstance(value.get(key), list):
            return value[key]
    return []


def main():
    report = []
    raw = OUTPUT.read_bytes()
    imported = requests.post(f"{BULK}/import", json={"filename": OUTPUT.name,
                             "content_base64": base64.b64encode(raw).decode()}, timeout=60).json()
    batch_id = imported["batch_id"]
    requests.post(f"{BULK}/import/{batch_id}/validate", timeout=60).raise_for_status()
    batch = requests.get(f"{BULK}/import/{batch_id}", timeout=60).json()["result"]
    products = batch["products"]
    report.append(("基础导入", "通过" if imported["errors"] == 0 else "失败",
                   f"7行预期/实际{imported['total']}行；3商品预期/实际{imported['product_count']}商品；错误{imported['errors']}"))
    report.append(("SKU合并", "通过" if [p["sku_count"] for p in products] == [4, 2, 1] else "失败",
                   str([(p["product_code"], p["sku_count"], p["stock_total"]) for p in products])))

    image_urls = sorted({url for product in products for url in product["main_images"] + product["detail_images"]})
    image_checks = []
    for url in image_urls:
        response = requests.get(url, timeout=30)
        image_checks.append(response.status_code == 200 and response.headers.get("content-type", "").startswith("image/"))
    report.append(("图片下载", "通过" if all(image_checks) else "失败", f"{sum(image_checks)}/{len(image_checks)} 张公网图片可下载"))

    alias = requests.get(f"{BULK}/category-aliases", timeout=30).json()["result"][0]
    wx_candidates = requests.get(f"{WECHAT}/categories", params={"keyword": "项链"}, timeout=180).json().get("results", [])
    latest = next(item for item in wx_candidates if item.get("leaf") and "天然钻石" in item.get("path", ""))
    saved_ids = [str(node.get("cat_id")) for node in alias["wechat"]["chain"]]
    latest_ids = [str(node.get("cat_id")) for node in latest["chain"]]
    report.append(("微信类目映射", "通过" if saved_ids == latest_ids else "警告",
                   f"已保存{saved_ids}；cats_v2最新{latest_ids}。不一致时必须更新映射后才能发布"))

    xhs_category_id = alias["xhs"]["category_id"]
    brands = result_list(requests.get(f"{XHS}/brands", params={"category_id": xhs_category_id, "keyword": ""}, timeout=60).json(), "brands")
    brand = next((item for item in brands if item.get("name") == "钻石世家"), None)
    report.append(("小红书品牌", "通过" if brand else "失败", f"品牌ID：{brand.get('id') if brand else '未找到'}"))

    shipping = result_list(requests.get(f"{XHS}/shipping-templates", timeout=60).json(), "carriageTemplateList", "templates")
    logistics = result_list(requests.get(f"{XHS}/logistics-plans", timeout=60).json(), "logisticsPlans", "list")
    report.append(("小红书店铺设置", "通过" if shipping and logistics else "失败",
                   f"运费模板{len(shipping)}个；有效物流方案{len([x for x in logistics if x.get('isValid') is not False])}个"))

    attrs_data = requests.get(f"{XHS}/category-attributes", params={"category_id": xhs_category_id}, timeout=60).json()
    attr_defs = result_list(attrs_data, "attributeV3s", "attributes")
    excel_attrs = products[0]["attributes"]
    xhs_attrs = {}
    required_total = len([item for item in attr_defs if item.get("isRequired")])
    required_matched = 0
    for definition in attr_defs:
        source_value = excel_attrs.get(definition.get("name"))
        if not source_value:
            continue
        values = result_list(requests.get(f"{XHS}/attribute-values", params={"category_id": xhs_category_id,
                             "attribute_id": definition.get("id")}, timeout=60).json(), "attributeValueV3s", "values")
        option = next((item for item in values if str(item.get("valueName", "")).lower() == str(source_value).lower()), None)
        if option:
            xhs_attrs[str(definition["id"])] = {"propertyId": definition["id"], "name": definition["name"],
                                                "valueId": option["valueId"], "value": option["valueName"]}
            if definition.get("isRequired"):
                required_matched += 1
    report.append(("小红书必填属性", "通过" if required_matched == required_total else "警告",
                   f"必填{required_total}项，Excel精确匹配{required_matched}项"))

    var_defs = result_list(requests.get(f"{XHS}/category-variations", params={"category_id": xhs_category_id}, timeout=60).json(), "variations")
    candidates = {}
    for definition in var_defs:
        candidates[str(definition["id"])] = result_list(requests.get(f"{XHS}/attribute-values", params={
            "category_id": xhs_category_id, "attribute_id": definition["id"]}, timeout=60).json(), "attributeValueV3s", "values")
    spec_map = {str(definition["id"]): definition["name"] for definition in var_defs if definition["name"] in products[0]["spec_dimensions"]}

    wx_mapping = {"wechat_category_chain": latest["chain"], "wechat_freight_template_id": "1057351904004",
                  "wechat_brand_id": "10002926", "wechat_attrs": excel_attrs}
    wx_payload = bulk_api._wechat_payload(products[0], wx_mapping)
    xhs_mapping = {"xhs_category_id": xhs_category_id, "xhs_brand_id": str(brand["id"]),
                   "xhs_shipping_template_id": str(shipping[0]["templateId"]),
                   "xhs_logistics_plan_id": str(logistics[0]["planInfoId"]), "xhs_attrs": xhs_attrs}
    xhs_payload = bulk_api._xhs_payload(products[0], xhs_mapping, {"xhs": {"var_defs": var_defs,
                                             "spec_map": spec_map, "candidates": candidates}})
    payload_ok = len(wx_payload["skus"]) == 4 and len(xhs_payload["sku_list"]) == 4 \
        and wx_payload["skus"][0]["sale_price"] == 799900 and xhs_payload["sku_list"][0]["price"] == 799900
    report.append(("发布请求转换", "通过" if payload_ok else "失败",
                   f"微信SKU {len(wx_payload['skus'])}个；小红书SKU {len(xhs_payload['sku_list'])}个；价格已由元转分"))
    report.append(("真实平台提交", "未执行", "为避免创建测试商品，本次未调用 /publish；图片也未写入平台素材库"))

    workbook = load_workbook(OUTPUT)
    if "测试记录" in workbook.sheetnames:
        del workbook["测试记录"]
    sheet = workbook.create_sheet("测试记录")
    sheet.append(["检查项", "结果", "实际情况"])
    for item in report:
        sheet.append(item)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 100
    workbook.save(OUTPUT)
    print(json.dumps({"file": str(OUTPUT), "batch_id": batch_id, "report": report}, ensure_ascii=True))


if __name__ == "__main__":
    main()
