# -*- coding: utf-8 -*-
"""按货盘真实编码生成「多规格 SKU」测试表。

直接复用 huopai_adapter 的货盘转换逻辑, 只保留指定款号, 因此
商品编码 / SKU编码 / 内部类目 / 商品属性 / 价格 / 库存 / 主图
与生产转换结果完全一致。

规格维度固定 2 个(货盘规则):
    规格1 = 主钻分数   值从货盘「规格」列解析(如 "主钻约50分" → "50分")
    规格2 = 戒托颜色   值从商家编码颜色字母提取(R=红色 W=白色 Y=黄色 S=银色)
发布时前端把规格名映射到小红书维度: 戒托颜色 → 颜色分类;
主钻分数 → 按类目实际存在的维度择优: 有「尺寸」挂尺寸, 否则挂「套装规格」(自由输入维度)。

编码规则(与货盘一致):
    商家编码(=SKU编码) = 款号主干 + 颜色字母 + 可选段 + "-" + 档位, 如 ZSTZ106XLW-10
    商品编码           = 去掉颜色字母和所有后缀的主干, 如 ZSTZ106XL

用法: python generate_multi_sku_test_excel.py
输出: 桌面 / 多规格SKU测试表.xlsx
"""

import collections
import os
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE_DIR)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import huopai_adapter as ha  # noqa: E402  复用生产转换逻辑

# 选中的货盘款号(商品编码主干): 组合唯一且不带 A/B 链型段
PICKS = ["ZSTZ106XL", "ZSTZ106SL", "ZNJ1725"]
PICK_SET = set(PICKS)

OUTPUT = os.path.join(os.path.expanduser("~"), "Desktop", "多规格SKU测试表.xlsx")

HEADERS = ha.HEADERS  # 与系统导入模板完全一致

# 货盘图片匹配不到时的兜底主图(1:1, >=1200px)
FALLBACK_IMAGE = ("https://images.unsplash.com/photo-1515562141207-7a88fb7ce338"
                  "?auto=format&fit=crop&w=1200&h=1200&q=90")


def load_rows():
    """跑一遍货盘转换, 拿到与生产一致的转换结果。"""
    images = ha.load_images()
    print("uploaded_images.json 图片数: %d" % len(images))
    workbook = openpyxl.load_workbook(ha.HUOPAI_PATH, read_only=True, data_only=True)
    rows = ha.parse_peiyuzuan(workbook["培育钻"], images)
    ha.enrich_split_chain(rows)
    ha.enrich_carat(rows)
    ha.enrich_titles(rows)   # 标题统一在最后按商品编码生成
    return rows


def describe(picked):
    """按商品编码汇总规格维度, 便于确认这一批能测出什么。"""
    by_code = collections.OrderedDict()
    for row in picked:
        by_code.setdefault(row["商品编码"], []).append(row)
    print("-" * 74)
    for code, items in by_code.items():
        dims = collections.OrderedDict()
        for item in items:
            for idx in (1, 2, 3):
                name = item.get("规格%d名称" % idx)
                val = item.get("规格%d值" % idx)
                if name:
                    dims.setdefault(name, collections.OrderedDict())
                    if val:
                        dims[name][val] = dims[name].get(val, 0) + 1
        combos = [(i.get("规格1值"), i.get("规格2值")) for i in items]
        dup = len(combos) - len(set(combos))
        print("%-16s %-26s SKU=%-3d 维度=%s 重复组合=%d" % (
            code, items[0]["内部类目"], len(items),
            [(k, list(v.keys())) for k, v in dims.items()], dup))
    return by_code


def build_workbook(picked, by_code):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品SKU"
    sheet.append(HEADERS)
    for row in picked:
        values = []
        for header in HEADERS:
            value = row.get(header, "")
            if header == "主图" and not value:
                value = FALLBACK_IMAGE
            values.append(value)
        sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    widths = [18, 30, 26, 30, 26, 12, 40, 90, 10, 46, 46, 22, 14, 14, 14, 14, 14, 14, 12, 12, 8]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(HEADERS)), sheet.max_row)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes = workbook.create_sheet("填写说明")
    notes.append(["项目", "说明"])
    notes.append(["数据来源", "全部来自货盘表经 huopai_adapter 的真实转换结果, 编码/类目/属性/价格/库存均一致。"])
    notes.append(["编码规则", "SKU编码 = 款号主干+颜色字母+可选段-档位(如 ZSTZ106XLW-10); 商品编码 = 去掉颜色字母和后缀的主干(如 ZSTZ106XL)。"])
    notes.append(["规格维度(固定2个)", "规格1 = 主钻分数(从货盘「规格」列解析); 规格2 = 戒托颜色(从编码颜色字母 W/Y/R/S 提取)。"])
    notes.append(["小红书维度映射", "戒托颜色 → 颜色分类; 主钻分数 → 有「尺寸」维度时挂尺寸, 否则挂「套装规格」。发布时由前端自动完成。"])
    notes.append(["为什么必须是这两个", "同款「钻石颜色」字段都是无色, 规格值会重复 → 平台只显示一个规格; 主钻分数+戒托颜色能保证组合唯一。"])
    notes.append(["商品合并", "商品编码相同的行合并为一个商品; 本次共 %d 个商品、%d 个 SKU。" % (len(by_code), len(picked))])
    notes.append(["微信侧限制", "珠宝类目除裸钻外 sale_attr_list 为空, 微信商品页不展示规格维度; 多规格主要看小红书。"])
    notes.append(["价格", "Excel 用元(货盘零售标价按 /0.7 换算); 发布请求自动转分。划线价按货盘规则不填。"])
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 118
    for row_cells in notes.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    expected = workbook.create_sheet("预期结果")
    expected.append(["商品编码", "内部类目", "SKU数", "主钻分数", "戒托颜色", "预期"])
    for code, items in by_code.items():
        scores, colors = [], []
        for item in items:
            score, color = item.get("规格1值") or "", item.get("规格2值") or ""
            if score and score not in scores:
                scores.append(score)
            if color and color not in colors:
                colors.append(color)
        expected.append([code, items[0]["内部类目"], len(items),
                         "/".join(scores), "/".join(colors),
                         "小红书应展示 2 个规格维度选择器, 共 %d 个 SKU 组合可切换" % len(items)])
    expected.append(["(全部)", "-", len(picked), "-", "-",
                     "所有 SKU 的(主钻分数 x 戒托颜色)组合唯一, 平台才能展示完整的多规格选择器"])
    for cell in expected[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
    for column, width in zip("ABCDEF", [18, 26, 8, 22, 26, 56]):
        expected.column_dimensions[column].width = width
    for row_cells in expected.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    workbook.save(OUTPUT)
    return OUTPUT


def main():
    rows = load_rows()
    picked = [r for r in rows if str(r.get("商品编码", "")).split("-")[0] in PICK_SET]
    picked.sort(key=lambda r: (str(r.get("商品编码")), str(r.get("SKU编码"))))
    if not picked:
        raise SystemExit("货盘里没找到指定款号: %s" % PICKS)
    by_code = describe(picked)
    path = build_workbook(picked, by_code)
    print("-" * 74)
    print("已生成: %s" % path)
    print("商品数: %d, SKU 行数: %d" % (len(by_code), len(picked)))


if __name__ == "__main__":
    main()
