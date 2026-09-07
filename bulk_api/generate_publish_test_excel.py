# -*- coding: utf-8 -*-
"""Generate a detailed, non-production Excel fixture for bulk publishing tests."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT = Path.home() / "Desktop" / "批量商品发布完整测试数据.xlsx"
HEADERS = [
    "商品编码", "标题", "微信标题", "小红书标题", "内部类目", "品牌", "描述", "商品属性", "重量",
    "主图", "详情图", "SKU编码", "规格1名称", "规格1值", "规格2名称", "规格2值", "规格3名称",
    "规格3值", "原价（元）", "售价（元）", "库存",
]

RING_IMAGE = "https://images.unsplash.com/photo-1605100804763-247f67b3557e?auto=format&fit=crop&w=1200&q=90"
NECKLACE_IMAGE = "https://images.unsplash.com/photo-1599643478518-a784e5dc4c8f?auto=format&fit=crop&w=1200&q=90"
DETAIL_IMAGE = "https://images.unsplash.com/photo-1515562141207-7a88fb7ce338?auto=format&fit=crop&w=1200&q=90"
CATEGORY = "钻石 > 钻石首饰 > 项链"
ATTRS = "钻石切工=Very good/优良;副钻分数=10分以下;形状=圆形;镶嵌材质=18K金;镶嵌方式=四爪镶;材质=18K金;颜色=无色;款式=项链"


def row(code, title, wx_title, xhs_title, description, weight, image, sku, color, clarity, carat,
        original_price, price, stock):
    return [
        code, title, wx_title, xhs_title, CATEGORY, "钻石世家", description, ATTRS, weight,
        image, DETAIL_IMAGE, sku, "颜色分类", color, "钻石净度", clarity, "重量/克拉", carat,
        original_price, price, stock,
    ]


ROWS = [
    row("TEST-NECKLACE-001", "18K金圆形钻石项链经典款", "18K金圆形钻石项链", "18K金圆形钻石项链经典锁骨链",
        "测试商品一：用于验证同一商品多SKU、三规格、价格库存和平台标题。", 8.6, NECKLACE_IMAGE,
        "TEST-N001-W-VS-030", "银色", "VS/微瑕", "0.30ct", 8999, 7999, 12),
    row("TEST-NECKLACE-001", "18K金圆形钻石项链经典款", "18K金圆形钻石项链", "18K金圆形钻石项链经典锁骨链",
        "测试商品一：用于验证同一商品多SKU、三规格、价格库存和平台标题。", 8.6, NECKLACE_IMAGE,
        "TEST-N001-W-SI-050", "银色", "SI/小瑕", "0.50ct", 12999, 10999, 8),
    row("TEST-NECKLACE-001", "18K金圆形钻石项链经典款", "18K金圆形钻石项链", "18K金圆形钻石项链经典锁骨链",
        "测试商品一：用于验证同一商品多SKU、三规格、价格库存和平台标题。", 8.6, NECKLACE_IMAGE,
        "TEST-N001-R-VS-030", "金色", "VS/微瑕", "0.30ct", 9299, 8299, 5),
    row("TEST-NECKLACE-001", "18K金圆形钻石项链经典款", "18K金圆形钻石项链", "18K金圆形钻石项链经典锁骨链",
        "测试商品一：用于验证同一商品多SKU、三规格、价格库存和平台标题。", 8.6, NECKLACE_IMAGE,
        "TEST-N001-R-SI-050", "金色", "SI/小瑕", "0.50ct", 13299, 11299, 3),
    row("TEST-NECKLACE-002", "18K金钻石吊坠简约通勤款", "18K金钻石吊坠通勤款", "18K金钻石吊坠简约通勤锁骨链",
        "测试商品二：用于验证双SKU、零库存SKU和详情图片。", 6.2, RING_IMAGE,
        "TEST-N002-W-VVS-020", "银色", "VVS/极微瑕", "0.20ct", 6999, 5999, 20),
    row("TEST-NECKLACE-002", "18K金钻石吊坠简约通勤款", "18K金钻石吊坠通勤款", "18K金钻石吊坠简约通勤锁骨链",
        "测试商品二：用于验证双SKU、零库存SKU和详情图片。", 6.2, RING_IMAGE,
        "TEST-N002-W-VS-030", "银色", "VS/微瑕", "0.30ct", 7999, 6899, 0),
    row("TEST-NECKLACE-003", "18K金钻石项链礼赠精选款", "18K金钻石项链礼赠款", "18K金钻石项链礼赠精选款",
        "测试商品三：用于验证单SKU商品。", 7.0, NECKLACE_IMAGE,
        "TEST-N003-W-VS-025", "银色", "VS/微瑕", "0.25ct", 7599, 6599, 30),
]


def main():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品SKU"
    sheet.append(HEADERS)
    for item in ROWS:
        sheet.append(item)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(ROWS) + 1}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [20, 30, 27, 34, 28, 14, 42, 68, 10, 48, 48, 24, 14, 14, 14, 14, 14, 14, 14, 14, 10]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    notes = workbook.create_sheet("填写说明")
    notes.append(["项目", "说明"])
    notes.append(["测试范围", "仅用于导入、合并、校验、类目属性映射、图片转存和发布请求转换测试。正式发布前请再次确认。"])
    notes.append(["图片", "主图和详情图是公网 URL。当前系统不读取 Excel 内嵌图片，发布时 Worker 会分别上传到微信和小红书素材服务器。"])
    notes.append(["商品合并", "商品编码相同的行应合并为一个商品；本表应合并为 3 个商品、7 个 SKU。"])
    notes.append(["类目", f"使用系统已有映射的内部类目：{CATEGORY}。"])
    notes.append(["规格", "覆盖颜色分类、钻石净度、重量/克拉三个规格维度。小红书必须将规格值匹配为平台 valueId。"])
    notes.append(["价格", "Excel 使用元；发布请求中微信和小红书应转换为分。"])
    notes.append(["安全", "所有商品编码和 SKU 编码都以 TEST- 开头，防止和正式货号混淆。"])
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 100
    for row_cells in notes.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    expected = workbook.create_sheet("预期结果")
    expected.append(["检查项", "预期"])
    expected.append(["导入行数", 7])
    expected.append(["合并商品数", 3])
    expected.append(["基础校验错误", 0])
    expected.append(["TEST-NECKLACE-001", "4个SKU，库存合计28，售价范围7999-11299元"])
    expected.append(["TEST-NECKLACE-002", "2个SKU，库存合计20，允许一个SKU库存为0"])
    expected.append(["TEST-NECKLACE-003", "1个SKU，库存30"])
    expected.append(["类目映射", "微信和小红书都应命中已保存的类目映射"])
    expected.append(["图片处理", "发布前各平台分别返回平台素材URL，相同源图片应命中缓存"])
    for cell in expected[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6B4F")
    expected.column_dimensions["A"].width = 28
    expected.column_dimensions["B"].width = 72

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
